import datetime, html, json
from openpyxl import load_workbook
wb=load_workbook('demo.xlsx',data_only=True)
ACC="#1F5E8C"
def fmt(v,kind=None):
    if v is None or v=="": return ""
    if isinstance(v,(datetime.datetime,datetime.date)): return v.strftime("%b %-d, %Y")
    if kind=="$": return ("-" if v==0 else (("−$" if v<0 else "$")+f"{abs(v):,.2f}"))
    if kind=="$0": return ("-" if v==0 else ("$%s" % f"{v:,.0f}"))
    if kind=="%": return f"{v*100:.0f}%"
    if isinstance(v,float): return f"{v:g}"
    return html.escape(str(v))
def table(ws,r0,r1,cols,kinds,inputs=(),hdr_row=4,status_col=None):
    h="<table><tr>"+"".join(f"<th>{html.escape(str(ws.cell(hdr_row,c).value))}</th>" for c in cols)+"</tr>"
    for r in range(r0,r1+1):
        h+="<tr>"
        for c,k in zip(cols,kinds):
            v=ws.cell(r,c).value; cls="in" if c in inputs else ""
            t=fmt(v,k)
            if status_col==c: cls+=" st-"+str(v).replace(" ","")
            if c==11 and ws.title=="Jobs" and v=="No": cls+=" st-Overdue"
            h+=f'<td class="{cls}">{t}</td>'
        h+="</tr>"
    return h+"</table>"
def win(tabname,body,w=1700):
    tabs=["Start Here","Dashboard","Clients","Jobs","Expenses","Quote Calculator","Invoice","Settings"]
    t="".join(f'<span class="tab {"on" if x==tabname else ""}">{x}</span>' for x in tabs)
    return f'<div class="win" style="width:{w}px"><div class="bar"><i></i><i></i><i></i><span>Pressure Washing Business Manager</span></div><div class="sheet">{body}</div><div class="tabs">{t}</div></div>'
CSS=f"""*{{box-sizing:border-box}}body{{margin:0;width:2000px;height:1600px;font-family:Inter,Arial,sans-serif;background:linear-gradient(135deg,#E7EFF7,#F8F5EE);display:flex;flex-direction:column;align-items:center;justify-content:center;overflow:hidden}}
h1{{font-size:78px;margin:0 0 10px;color:#15324a;letter-spacing:-1px}} h2{{font-size:38px;margin:0 0 40px;color:{ACC};font-weight:600}}
.win{{background:#fff;border-radius:18px;box-shadow:0 30px 80px rgba(20,50,50,.25);overflow:hidden}}
.bar{{background:#f1f3f4;height:48px;display:flex;align-items:center;padding:0 18px;gap:9px;font-size:20px;color:#555}}.bar i{{width:14px;height:14px;border-radius:50%;background:#e0e0e0;display:inline-block}}.bar span{{margin-left:16px}}
.sheet{{padding:26px 30px}} .tabs{{background:#f1f3f4;padding:10px 20px;display:flex;gap:6px}}.tab{{font-size:17px;padding:7px 14px;border-radius:6px;color:#555}}.tab.on{{background:#fff;color:{ACC};font-weight:700;box-shadow:0 1px 3px rgba(0,0,0,.15)}}
table{{border-collapse:collapse;width:100%;font-size:20px}}th{{background:{ACC};color:#fff;padding:11px 10px;font-weight:600;text-align:left}}td{{padding:9px 10px;border-bottom:1px solid #e3e8ea;white-space:nowrap}}td.in{{background:#FFF7D6;color:#1F4EB4}}
.st-Overdue{{background:#F8D7DA!important;color:#9C1C28!important;font-weight:700}}.st-Duesoon{{background:#FFF3CD!important;font-weight:600}}
.kpis{{display:flex;gap:22px;margin-bottom:24px}}.kpi{{flex:1;background:#E7EFF7;border-radius:14px;padding:18px 22px}}.kpi b{{display:block;font-size:44px;color:{ACC}}}.kpi span{{font-size:20px;color:#57606a}}
.ttl{{font-size:34px;font-weight:800;color:{ACC};margin-bottom:6px}}.sub{{font-size:19px;color:#57606a;margin-bottom:20px;font-style:italic}}
.badge{{position:absolute;top:60px;right:70px;background:{ACC};color:#fff;font-size:30px;font-weight:700;padding:16px 28px;border-radius:40px}}
"""
def page(name,inner,head=None,sub=None):
    hh=(f"<h1>{head}</h1>" if head else "")+(f"<h2>{sub}</h2>" if sub else "")
    open(f"img_{name}.html","w").write(f"<html><head><style>{CSS}</style></head><body>{hh}{inner}</body></html>")
d=wb['Dashboard']
kp="".join(f'<div class="kpi"><span>{d.cell(5,c).value}</span><b>{fmt(d.cell(6,c).value,"$0")}</b></div>' for c in (2,4,6,8))
rows=[r for r in range(9,21) if (d.cell(r,2).value or 0)>0]
bars=""; mx=max(d.cell(r,2).value for r in rows)
for r in rows:
    rv=d.cell(r,2).value; pf=d.cell(r,6).value
    bars+=f'<div style="display:flex;flex-direction:column;align-items:center;gap:6px;flex:1"><div style="display:flex;gap:8px;align-items:flex-end;height:260px"><div style="width:46px;height:{260*rv/mx}px;background:{ACC};border-radius:6px 6px 0 0"></div><div style="width:46px;height:{max(0,260*pf/mx)}px;background:#93BDE0;border-radius:6px 6px 0 0"></div></div><span style="font-size:18px;color:#555">{d.cell(r,1).value}</span></div>'
chart=f'<div style="display:flex;gap:20px;padding:10px 0 0">{bars}</div><div style="font-size:18px;color:#555;margin-top:8px"><span style="color:{ACC}">■</span> Revenue &nbsp; <span style="color:#93BDE0">■</span> Profit</div>'
dash=f'<div class="ttl">Dashboard — {d["B4"].value}</div><div class="sub">Updates automatically from your Jobs and Expenses</div><div class="kpis">{kp}</div><div style="display:flex;gap:30px"><div style="flex:1.2">{table(d,9,21,range(1,9),[None,"$0","$0","$0",None,"$0","%","$0"],hdr_row=8).replace("<tr><td class=\"\">Jan","<tr><td>Jan")}</div><div style="flex:1">{chart}</div></div>'
# filter months w/o data from table for compactness
K=[None,"$0","$0","$0",None,"$0","%","$0"]
mt='<table><tr>'+"".join(f"<th>{d.cell(8,c).value}</th>" for c in range(1,9))+'</tr>'+"".join('<tr>'+"".join(f'<td>{fmt(d.cell(r,c).value,k)}</td>' for c,k in zip(range(1,9),K))+'</tr>' for r in rows+[21])+'</table>'
hero=f'<div class="ttl">Dashboard — {d["B4"].value}</div><div class="sub">Updates automatically from your Jobs and Expenses</div><div class="kpis">{kp}</div><div style="display:flex;gap:30px"><div style="flex:1.2">{mt}</div><div style="flex:1">{chart}</div></div>'
page("1_hero",win("Dashboard",hero,1800),"Pressure Washing Business Manager","Clients · Jobs · Quotes · Invoices · Profit — Google Sheets & Excel")
page("2_dashboard",win("Dashboard",dash,1880),None,None)
q=wb['Quote Calculator']
qi="".join(f'<tr><td>{q.cell(r,1).value}</td><td class="in">{fmt(q.cell(r,2).value)}</td></tr>' for r in range(4,9))
qa="".join(f'<tr><td>{q.cell(r,4).value}</td><td class="in">{q.cell(r,5).value}</td></tr>' for r in range(4,10))
qo="".join(f'<tr style="{"font-weight:800;font-size:28px;color:"+ACC if q.cell(r,1).value=="QUOTE PRICE" else ""}"><td>{q.cell(r,1).value}</td><td>{fmt(q.cell(r,2).value,"$" if r!=21 else None)}</td></tr>' for r in list(range(12,20))+[21,22,23,24])
qb=f'<div class="ttl">Quote Calculator</div><div class="sub">Price any job in under a minute — house size, stories, extra surfaces, add-ons and frequency</div><div style="display:flex;gap:40px"><div style="flex:1"><table><tr><th colspan=2>Property details</th></tr>{qi}</table><br><table><tr><th colspan=2>Add-ons</th></tr>{qa}</table></div><div style="flex:1"><table><tr><th colspan=2>Your quote</th></tr>{qo}</table></div></div>'
page("3_quote",win("Quote Calculator",qb,1600),"Quote with confidence","Know your price, hours and profit before you say yes")
c=wb['Clients']
cb=f'<div class="ttl">Clients</div><div class="sub">Next due date and overdue alerts fill in automatically</div>'+table(c,5,13,[1,5,6,7,10,11,12,13,14],[None,None,None,"$",None,"$",None,None,None],inputs=(1,5,6,7),status_col=14)
page("4_clients",win("Clients",cb,1880),"Never miss a repeat customer","Annual, 6-month and quarterly clients come back on schedule")
j=wb['Jobs']
last=max(r for r in range(5,1005) if j.cell(r,3).value)
jbb=f'<div class="ttl">Jobs</div><div class="sub">Pick a client — price fills in. Profit per job after labor, supplies and mileage.</div>'+table(j,last-11,last,[1,2,3,5,6,10,11,12,15,17],[None,None,None,None,None,"$",None,None,"$","$"],inputs=(2,3,5,6,11,12))
page("5_jobs",win("Jobs",jbb,1880),"Profit on every visit","Crew labor, chemicals and mileage subtracted automatically")
i=wb['Invoice']
inv=f'''<div style="width:1000px;padding:20px 30px;font-size:22px"><div style="display:flex;justify-content:space-between"><div><b style="font-size:30px">{i["A6"].value}</b><br>{i["A7"].value}<br>{i["A8"].value}<br>{i["A9"].value}</div><div style="text-align:right"><b style="font-size:46px;color:{ACC}">INVOICE</b><br>Invoice # {i["E6"].value}<br>Service date {fmt(i["E7"].value)}<br>Due {fmt(i["E8"].value)}<br><b style="color:#9C1C28">{i["E9"].value}</b></div></div><br><b style="color:{ACC}">BILL TO</b><br><b>{i["A12"].value}</b><br>{i["A13"].value}<br>{i["A14"].value}<br><br><table><tr><th>Description</th><th style="text-align:right">Amount</th></tr><tr><td>{i["A17"].value}</td><td style="text-align:right">{fmt(i["E17"].value,"$")}</td></tr><tr><td>Add-on services</td><td style="text-align:right">{fmt(i["E18"].value,"$")}</td></tr><tr><td style="text-align:right"><b>TOTAL</b></td><td style="text-align:right"><b style="font-size:28px">{fmt(i["E20"].value,"$")}</b></td></tr></table><br><i style="color:{ACC}">Thank you for your business!</i><br>{i["A24"].value}</div>'''
page("6_invoice",win("Invoice",f'<div class="sub">Type a Job ID → the invoice fills itself in. Print or save as PDF.</div><div style="display:flex;justify-content:center">{inv}</div>',1300),"One-click invoices",None)
feat=[("📊","Profit dashboard","Revenue, costs, profit, margin and unpaid balance by month"),("🧮","Quote calculator","House size, stories, surfaces, add-ons, discounts, hours & profit"),("👥","Client tracker","Annual, 6-month and quarterly repeat dates with overdue alerts"),("🧾","Jobs log","Price auto-fills; labor, supplies & mileage subtracted per job"),("📄","Invoice generator","Enter a Job ID, print or save as PDF"),("💸","Expense tracker","Categories, receipts, and tax set-aside estimate"),("⚙️","Fully editable","Your services, prices, rates, crew and categories"),("✅","Google Sheets + Excel","Instant download, works in both — no subscription")]
fb="".join(f'<div style="background:#fff;border-radius:16px;padding:28px;box-shadow:0 10px 30px rgba(20,50,50,.12)"><div style="font-size:46px">{a}</div><div style="font-size:30px;font-weight:800;color:#15324a;margin:8px 0">{b}</div><div style="font-size:22px;color:#57606a">{c}</div></div>' for a,b,c in feat)
page("7_included",f'<div style="display:grid;grid-template-columns:repeat(4,400px);gap:26px">{fb}</div>',"What's included","8 connected tabs · instant digital download")
print("ok")
