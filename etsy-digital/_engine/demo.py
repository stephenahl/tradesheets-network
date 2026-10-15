import random
from datetime import date,timedelta
from openpyxl import load_workbook
random.seed(7)
wb=load_workbook('Pressure-Washing-Business-Manager.xlsx')
cl=wb['Clients']; jb=wb['Jobs']; xp=wb['Expenses']
clients=[("Johnson Family","House Soft Wash","Annual",350),("Bright Dental","Commercial Flatwork","Quarterly",300),("Patel Residence","Driveway & Walkways","Every 6 months",160),
("Lakeview HOA","Commercial Flatwork","Monthly",420),("Garcia Home","House Soft Wash","Annual",325),("Oak Street Cafe","Commercial Flatwork","Monthly",260),
("Miller Family","Roof Soft Wash","One-time",475),("Summit Law Office","Commercial Flatwork","Quarterly",280),("Rivera Home","Driveway & Walkways","Every 6 months",150),
("Chen Residence","Deck / Patio Clean","Annual",240),("Nguyen Residence","House Soft Wash","One-time",340),("Baker Home","Fence Cleaning","One-time",210)]+[(f"{n} Residence",random.choice(["House Soft Wash","Driveway & Walkways","Deck / Patio Clean","Roof Soft Wash","Fence Cleaning"]),"One-time",random.choice([175,225,295,340,425])) for n in ["Adams","Brooks","Carter","Diaz","Evans","Foster","Gray","Hughes","Irwin","Jensen","Kim","Lopez","Morgan","Ortiz","Price","Reed"]]
streets=[f"{random.randint(10,999)} {random.choice(['Oak','Pine','Elm','Birch','Cedar','Maple','Aspen'])} St" for _ in range(40)]
step={"Annual":365,"Every 6 months":182,"Quarterly":91,"Monthly":30}
hrs={"House Soft Wash":3,"Commercial Flatwork":3,"Driveway & Walkways":1.5,"Deck / Patio Clean":2.5,"Roof Soft Wash":4,"Fence Cleaning":2.5}
jobs=[]
for i,(n,s,f,p) in enumerate(clients):
    r=5+i
    for j,v in enumerate([n,f"(555) 0{random.randint(10,99)}-{random.randint(1000,9999)}",n.split()[0].lower()+"@example.com",streets[i],s,f,p,date(2026,5,1),""]): cl.cell(r,1+j,v)
    d=date(2026,6,1)+timedelta(days=random.randint(0,40 if f!='One-time' else 130))
    end=date(2026,9,25) if n=='Oak Street Cafe' else date(2026,10,13)
    while d<=end:
        jobs.append((d,n,s,p))
        if f=="One-time" or f not in step: break
        d+=timedelta(days=step[f])
# extra one-time jobs from walk-ins (as one-time clients above)
jobs.sort()
for k,(d,n,s,p) in enumerate(jobs):
    r=5+k; paid="No" if d>date(2026,10,1) and random.random()<0.6 else "Yes"
    vals={2:d,3:n,5:hrs[s]+random.choice([0,0,0.5]),6:random.choice(["Me","Crew 2"]),9:random.choice([0,0,50,75,85]),11:paid,12:None if paid=="No" else random.choice(["Zelle","Card","Check"]),13:round(random.uniform(12,30),2),14:random.randint(4,20)}
    for c,v in vals.items(): jb.cell(r,c,v)
    for c in (4,7,18): jb.cell(r,c,None)
ex=[]
for m in range(6,11):
    ex+=[(date(2026,m,3),"Chemicals & detergents","Pool supply store",round(random.uniform(90,160),2)),(date(2026,m,5),"Insurance","Next Insurance",65),(date(2026,m,12),"Fuel","Shell",round(random.uniform(80,140),2)),(date(2026,m,22),"Equipment & repairs","Pressure Tek",round(random.uniform(20,120),2))]
for k,(d,c,v,a) in enumerate(ex):
    for j,val in enumerate([d,c,v,"",a,"Card","Yes"]): xp.cell(5+k,1+j,val)
wb['Invoice']['B4']="J-%04d"%(len(jobs)-1)
wb.save('demo.xlsx'); print(len(jobs),"jobs")
