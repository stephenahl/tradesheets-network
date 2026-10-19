import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from datetime import date

T = dict(name="Lawn Care Business Manager", trade="lawn care", accent="3E6B2F", light="EAF2E4")
if len(sys.argv) > 1: pass
ACC=T["accent"]; LIGHT=T["light"]
F="Arial"
def font(**k): k.setdefault("name",F); k.setdefault("size",10); return Font(**k)
HDR_FILL=PatternFill("solid",fgColor=ACC); IN_FILL=PatternFill("solid",fgColor="FFF7D6"); BAND=PatternFill("solid",fgColor=LIGHT)
thin=Side(style="thin",color="D0D7DE"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
BLUE="1F4EB4"
MONEY='$#,##0.00;($#,##0.00);"-"'; MONEY0='$#,##0;($#,##0);"-"'; DATE='mmm d, yyyy'; PCT='0.0%'
N=1000  # data rows

wb=Workbook()
def sheet(title, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title=title; ws.sheet_view.showGridLines=False
    ws.sheet_properties.tabColor=ACC
    return ws
def title(ws, text, sub, width_to="H"):
    ws["A1"]=text; ws["A1"].font=font(size=18,bold=True,color=ACC)
    ws["A2"]=sub; ws["A2"].font=font(size=10,italic=True,color="57606A")
    ws.row_dimensions[1].height=30
def header(ws,row,cols,widths=None):
    for i,c in enumerate(cols,1):
        cell=ws.cell(row=row,column=i,value=c)
        cell.font=font(bold=True,color="FFFFFF"); cell.fill=HDR_FILL
        cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=BOX
    ws.row_dimensions[row].height=32
    if widths:
        for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i) if i<=26 else 'A'+chr(64+i-26)].width=w
def inp(c,fmt=None):
    c.font=font(color=BLUE); c.fill=IN_FILL; c.border=BOX
    if fmt: c.number_format=fmt
def calc(c,fmt=None,bold=False):
    c.font=font(bold=bold); c.border=BOX
    if fmt: c.number_format=fmt

# ---------------- START HERE
s=sheet("Start Here",first=True)
title(s,T["name"],"Clients · Jobs · Quotes · Invoices · Expenses · Profit dashboard — Google Sheets & Excel")
s.column_dimensions["A"].width=4; s.column_dimensions["B"].width=34; s.column_dimensions["C"].width=70
steps=[("1. Settings","Enter your business name, hourly labor cost, tax set-aside %, mileage rate and your service prices."),
("2. Clients","Add each client once: name, contact, service, how often you service them, and their price. Next due date and overdue flags fill in automatically."),
("3. Jobs","Log every visit. Pick the client from the dropdown — price fills in from their record (type in 'Price override' to change it for one job). Mark Paid = Yes when paid."),
("4. Expenses","Log supplies, gas, equipment, insurance, ads. Categories come from Settings."),
("5. Quote Calculator","Price a new property in seconds from lawn size, beds, trees, add-ons and frequency. Shows estimated hours and your profit."),
("6. Invoice","Type a Job ID (e.g. J-0001). The invoice fills in. Print or File > Download as PDF."),
("7. Dashboard","Pick a year. See revenue, costs, profit, unpaid balance and tax set-aside by month.")]
s["B4"]="HOW TO USE"; s["B4"].font=font(bold=True,size=12,color=ACC)
for i,(a,b) in enumerate(steps):
    r=5+i; s.cell(r,2,a).font=font(bold=True); c=s.cell(r,3,b); c.font=font(); c.alignment=Alignment(wrap_text=True,vertical="top"); s.row_dimensions[r].height=30
r=5+len(steps)+1
s.cell(r,2,"COLOR KEY").font=font(bold=True,size=12,color=ACC)
k=s.cell(r+1,2,"Yellow cell, blue text"); inp(k); s.cell(r+1,3,"Type here — your inputs").font=font()
k=s.cell(r+2,2,"White cell, black text"); calc(k); s.cell(r+2,3,"Calculated automatically — don't type over these").font=font()
s.cell(r+4,2,"GOOGLE SHEETS").font=font(bold=True,size=12,color=ACC)
c=s.cell(r+5,3,"Upload this file to Google Drive, open it, then File > Save as Google Sheets. Everything works in both Excel and Google Sheets."); c.font=font(); c.alignment=Alignment(wrap_text=True); s.row_dimensions[r+5].height=30
s.cell(r+6,3,"The example rows (marked EXAMPLE) show the format — delete them when you start. After that, clear a job row instead of deleting it, so Job IDs and invoice numbers stay the same.").font=font(italic=True,color="57606A")

# ---------------- SETTINGS
st=sheet("Settings")
title(st,"Settings","Your business details and defaults. Every yellow cell is editable.")
st.column_dimensions["A"].width=34; st.column_dimensions["B"].width=18; st.column_dimensions["C"].width=4
st.column_dimensions["D"].width=26; st.column_dimensions["E"].width=14; st.column_dimensions["F"].width=14; st.column_dimensions["G"].width=4; st.column_dimensions["H"].width=24
rows=[("Business name","Green Edge Lawn Care",None),("Owner / contact","Your Name",None),("Phone","(555) 555-0123",None),("Email","hello@example.com",None),
("Hourly labor cost (what you pay per crew-hour)",20,MONEY),("Tax set-aside % of profit",0.25,PCT),("Mileage rate per mile",0.70,MONEY),("Invoice due in (days)",14,"0")]
for i,(a,b,f) in enumerate(rows):
    r=4+i; st.cell(r,1,a).font=font(); c=st.cell(r,2,b); inp(c,f)
st["B8"].comment=Comment("Pay rate plus payroll costs per crew-hour. Used for labor cost and profit.","Template")
st["B9"].comment=Comment("A planning estimate only. Ask a tax professional what to set aside.","Template")
st["B10"].comment=Comment("Example default. Check the current IRS standard mileage rate for your tax year.","Template")
# Services
st["D3"]="SERVICES"; st["D3"].font=font(bold=True,color=ACC)
for j,h in enumerate(["Service","Base price","Est. hours"]):
    c=st.cell(4,4+j,h); c.font=font(bold=True,color="FFFFFF"); c.fill=HDR_FILL; c.border=BOX
svcs=[("Weekly Mow",45,1.0),("Mow + Edge + Blow",55,1.25),("Spring Cleanup",180,4),("Fall Leaf Cleanup",220,5),("Aeration",120,1.5),("Commercial Mow",95,2),("",None,None),("",None,None)]
for i,(a,b,cc) in enumerate(svcs):
    r=5+i; inp(st.cell(r,4,a or None)); inp(st.cell(r,5,b),MONEY); inp(st.cell(r,6,cc),"0.0")
# Quote factors
st["D14"]="QUOTE FACTORS"; st["D14"].font=font(bold=True,color=ACC)
qf=[("Per extra flower bed (over 2)",6,MONEY),("Per extra tree/obstacle (over 2)",3,MONEY),("Per 1,000 sq ft over 5,000",5,MONEY),("Extra hours per flower bed",0.15,"0.00"),("Extra hours per tree/obstacle",0.05,"0.00"),("Extra hours per 1,000 sq ft over 5,000",0.1,"0.00")]
for i,(a,b,f) in enumerate(qf):
    r=15+i; st.cell(r,4,a).font=font(); inp(st.cell(r,5,b),f)
# Add-ons
st["D22"]="ADD-ONS"; st["D22"].font=font(bold=True,color=ACC)
for j,h in enumerate(["Add-on","Price","Hours"]):
    c=st.cell(23,4+j,h); c.font=font(bold=True,color="FFFFFF"); c.fill=HDR_FILL; c.border=BOX
adds=[("Hedge trimming",40,1),("Weed pulling",35,1),("Fertilizer application",45,0.5),("Mulch install",60,1),("Gutter cleaning",75,1.5),("Debris haul-away",40,0.5)]
for i,(a,b,cc) in enumerate(adds):
    r=24+i; inp(st.cell(r,4,a)); inp(st.cell(r,5,b),MONEY); inp(st.cell(r,6,cc),"0.00")
# Frequency
st["D31"]="FREQUENCY"; st["D31"].font=font(bold=True,color=ACC)
for j,h in enumerate(["Frequency","Days between","Discount"]):
    c=st.cell(32,4+j,h); c.font=font(bold=True,color="FFFFFF"); c.fill=HDR_FILL; c.border=BOX
fq=[("Weekly",7,0.10),("Every 2 weeks",14,0.05),("Every 4 weeks",28,0),("Monthly",30,0),("One-time",None,0)]
for i,(a,b,cc) in enumerate(fq):
    r=33+i; inp(st.cell(r,4,a)); inp(st.cell(r,5,b),"0"); inp(st.cell(r,6,cc),PCT)
# Lists
st["H3"]="EXPENSE CATEGORIES"; st["H3"].font=font(bold=True,color=ACC)
for i,a in enumerate(["Fuel","Equipment & repairs","Fertilizer & seed","Mulch & materials","Dump fees","Insurance","Advertising","Trailer / vehicle","Subcontractors","Other",""]):
    inp(st.cell(4+i,8,a or None))
st["H17"]="PAYMENT METHODS"; st["H17"].font=font(bold=True,color=ACC)
for i,a in enumerate(["Cash","Card","Check","Zelle","Venmo","Bank transfer",""]):
    inp(st.cell(18+i,8,a or None))
st["H27"]="CREW / STAFF"; st["H27"].font=font(bold=True,color=ACC)
for i,a in enumerate(["Me","Crew 2","","","",""]):
    inp(st.cell(28+i,8,a or None))

SVC="Settings!$D$5:$D$12"; SVCP="Settings!$E$5:$E$12"; SVCH="Settings!$F$5:$F$12"
FRQ="Settings!$D$33:$D$37"; FRQD="Settings!$E$33:$E$37"

# ---------------- CLIENTS
cl=sheet("Clients")
title(cl,"Clients","One row per client. Next due date and status update automatically from the Jobs log.")
cols=["Client name","Phone","Email","Address","Service","Frequency","Price per visit","Start date","Notes","Visits","Lifetime revenue","Last visit","Next due","Status"]
header(cl,4,cols,[24,15,24,30,18,15,13,13,26,8,14,13,13,13])
H=5
ex=[("EXAMPLE – Johnson Family","(555) 201-3344","jj@example.com","12 Oak St","Mow + Edge + Blow","Weekly",55,date(2026,9,1),"Gate code 1234 – dog in yard"),
("EXAMPLE – Bright Dental","(555) 870-1100","office@example.com","400 Main St","Commercial Mow","Every 2 weeks",95,date(2026,8,15),"Before 8am only")]
for r in range(H,H+300):
    e=ex[r-H] if r-H<len(ex) else [None]*9
    for j in range(9):
        c=cl.cell(r,1+j,e[j]); inp(c, MONEY if j==6 else DATE if j==7 else None)
    A=f"$A{r}"
    calc(cl.cell(r,10,f'=IF({A}="","",COUNTIFS(Jobs!$C$5:$C${4+N},{A}))'),"0")
    calc(cl.cell(r,11,f'=IF({A}="","",SUMIFS(Jobs!$J$5:$J${4+N},Jobs!$C$5:$C${4+N},{A}))'),MONEY)
    calc(cl.cell(r,12,f'=IF(OR({A}="",J{r}=0),"",_xlfn.MAXIFS(Jobs!$B$5:$B${4+N},Jobs!$C$5:$C${4+N},{A}))'),DATE)
    calc(cl.cell(r,13,f'=IF(OR({A}="",L{r}=""),"",IFERROR(IF(INDEX({FRQD},MATCH(F{r},{FRQ},0))="","",L{r}+INDEX({FRQD},MATCH(F{r},{FRQ},0))),""))'),DATE)
    calc(cl.cell(r,14,f'=IF({A}="","",IF(M{r}="",IF(J{r}=0,"New","—"),IF(M{r}<TODAY(),"Overdue",IF(M{r}-TODAY()<=3,"Due soon","Scheduled"))))'))
dv=DataValidation(type="list",formula1=SVC,allow_blank=True); cl.add_data_validation(dv); dv.add(f"E{H}:E{H+299}")
dv=DataValidation(type="list",formula1=FRQ,allow_blank=True); cl.add_data_validation(dv); dv.add(f"F{H}:F{H+299}")
cl.conditional_formatting.add(f"N{H}:N{H+299}",CellIsRule(operator="equal",formula=['"Overdue"'],fill=PatternFill("solid",fgColor="F8D7DA"),font=Font(name=F,color="9C1C28",bold=True)))
cl.conditional_formatting.add(f"N{H}:N{H+299}",CellIsRule(operator="equal",formula=['"Due soon"'],fill=PatternFill("solid",fgColor="FFF3CD")))
cl.freeze_panes="B5"; cl.auto_filter.ref=f"A4:N{H+299}"

# ---------------- JOBS
jb=sheet("Jobs")
title(jb,"Jobs","Log every visit. Price fills in from the client; type a Price override to change one job.")
cols=["Job ID","Date","Client","Service","Hours","Crew","Price override","Price","Add-on charges","Total charged","Paid?","Payment method","Supplies cost","Miles","Labor cost","Mileage cost","Profit","Notes"]
header(jb,4,cols,[9,13,24,17,8,12,11,11,11,12,8,13,11,8,11,11,12,24])
jex=[(date(2026,10,1),"EXAMPLE – Johnson Family",None,1.25,"Me",None,0,"Yes","Zelle",4,8),
(date(2026,10,2),"EXAMPLE – Bright Dental",None,2,"Crew 2",None,0,"Yes","Card",6,4),
(date(2026,10,9),"EXAMPLE – Bright Dental",None,2.5,"Crew 2",None,40,"No",None,6,4)]
for i in range(N):
    r=5+i; e=jex[i] if i<len(jex) else [None]*11
    calc(jb.cell(r,1,f'=IF(B{r}="","","J-"&TEXT(ROW()-4,"0000"))'))
    inp(jb.cell(r,2,e[0]),DATE); inp(jb.cell(r,3,e[1]))
    inp(jb.cell(r,4,e[2]))
    # service default from client if blank -> show via formula? keep input; price from client
    inp(jb.cell(r,5,e[3]),"0.0"); inp(jb.cell(r,6,e[4])); inp(jb.cell(r,7,e[5]),MONEY)
    calc(jb.cell(r,8,f'=IF(C{r}="","",IF(G{r}<>"",G{r},IFERROR(INDEX(Clients!$G$5:$G$304,MATCH(C{r},Clients!$A$5:$A$304,0)),0)))'),MONEY)
    inp(jb.cell(r,9,e[6]),MONEY)
    calc(jb.cell(r,10,f'=IF(C{r}="","",H{r}+N(I{r}))'),MONEY)
    inp(jb.cell(r,11,e[7])); inp(jb.cell(r,12,e[8])); inp(jb.cell(r,13,e[9]),MONEY); inp(jb.cell(r,14,e[10]),"0")
    calc(jb.cell(r,15,f'=IF(C{r}="","",N(E{r})*Settings!$B$8)'),MONEY)
    calc(jb.cell(r,16,f'=IF(C{r}="","",N(N{r})*Settings!$B$10)'),MONEY)
    calc(jb.cell(r,17,f'=IF(C{r}="","",J{r}-N(M{r})-O{r}-P{r})'),MONEY)
    inp(jb.cell(r,18))
    if i%2: 
        for cc in (1,8,10,15,16,17): jb.cell(r,cc).fill=BAND
L=f"{5}:{4+N}"
for rng,f in [(f"C5:C{4+N}","Clients!$A$5:$A$304"),(f"D5:D{4+N}",SVC),(f"F5:F{4+N}","Settings!$H$28:$H$33"),(f"K5:K{4+N}",'"Yes,No"'),(f"L5:L{4+N}","Settings!$H$18:$H$24")]:
    dv=DataValidation(type="list",formula1=f,allow_blank=True); jb.add_data_validation(dv); dv.add(rng)
jb.conditional_formatting.add(f"K5:K{4+N}",CellIsRule(operator="equal",formula=['"No"'],fill=PatternFill("solid",fgColor="F8D7DA"),font=Font(name=F,color="9C1C28",bold=True)))
jb.conditional_formatting.add(f"Q5:Q{4+N}",CellIsRule(operator="lessThan",formula=['0'],font=Font(name=F,color="9C1C28",bold=True)))
jb.freeze_panes="D5"; jb.auto_filter.ref=f"A4:R{4+N}"
JD=f"Jobs!$B$5:$B${4+N}"; JT=f"Jobs!$J$5:$J${4+N}"; JP=f"Jobs!$K$5:$K${4+N}"

# ---------------- EXPENSES
xp=sheet("Expenses")
title(xp,"Expenses","Every business cost that isn't already on a job (supplies on a job go in the Jobs tab).")
header(xp,4,["Date","Category","Vendor","Description","Amount","Paid with","Receipt saved?"],[13,20,22,32,12,15,14])
xex=[(date(2026,10,3),"Fuel","Shell","Mower and truck fuel",64.20,"Card","Yes"),(date(2026,10,5),"Insurance","Next Insurance","Monthly liability policy",48,"Card","Yes")]
xex=[(e[0],e[1],e[2],"EXAMPLE – "+e[3])+tuple(e[4:]) for e in xex]
for i in range(N):
    r=5+i; e=xex[i] if i<len(xex) else [None]*7
    for j in range(7): inp(xp.cell(r,1+j,e[j]),DATE if j==0 else MONEY if j==4 else None)
for rng,f in [(f"B5:B{4+N}","Settings!$H$4:$H$14"),(f"F5:F{4+N}","Settings!$H$18:$H$24"),(f"G5:G{4+N}",'"Yes,No"')]:
    dv=DataValidation(type="list",formula1=f,allow_blank=True); xp.add_data_validation(dv); dv.add(rng)
xp.freeze_panes="A5"; xp.auto_filter.ref=f"A4:G{4+N}"
XD=f"Expenses!$A$5:$A${4+N}"; XA=f"Expenses!$E$5:$E${4+N}"

# ---------------- QUOTE
q=sheet("Quote Calculator")
title(q,"Quote Calculator","Price a new client in under a minute. Change the yellow cells.")
q.column_dimensions["A"].width=34; q.column_dimensions["B"].width=20; q.column_dimensions["C"].width=4; q.column_dimensions["D"].width=24; q.column_dimensions["E"].width=12
qi=[("Service","Mow + Edge + Blow"),("Lawn area (sq ft)",9000),("Flower beds",4),("Trees / obstacles",5),("Frequency","Weekly")]
for i,(a,b) in enumerate(qi):
    r=4+i; q.cell(r,1,a).font=font(); inp(q.cell(r,2,b),"#,##0" if isinstance(b,int) else None)
dv=DataValidation(type="list",formula1=SVC); q.add_data_validation(dv); dv.add("B4")
dv=DataValidation(type="list",formula1=FRQ); q.add_data_validation(dv); dv.add("B8")
q["D3"]="Add-ons (Yes/No)"; q["D3"].font=font(bold=True,color=ACC)
for i in range(6):
    r=4+i; calc(q.cell(r,4,f'=IF(Settings!D{24+i}="","",Settings!D{24+i})')); inp(q.cell(r,5,"Yes" if i in (1,) else "No"))
dv=DataValidation(type="list",formula1='"Yes,No"'); q.add_data_validation(dv); dv.add("E4:E9")
out=[("Base price",f'=IFERROR(INDEX({SVCP},MATCH(B4,{SVC},0)),0)',MONEY),
("Flower bed adjustment",'=MAX(0,B6-2)*Settings!$E$15',MONEY),
("Tree/obstacle adjustment",'=MAX(0,B7-2)*Settings!$E$16',MONEY),
("Size adjustment",'=MAX(0,ROUNDUP((B5-5000)/1000,0))*Settings!$E$17',MONEY),
("Add-ons",'=SUMIFS(Settings!$E$24:$E$29,Settings!$D$24:$D$29,"<>",E4:E9,"Yes")',MONEY),
("Subtotal",'=SUM(B12:B16)',MONEY),
("Frequency discount",f'=-B17*IFERROR(INDEX(Settings!$F$33:$F$37,MATCH(B8,{FRQ},0)),0)',MONEY),
("QUOTE PRICE",'=ROUND(B17+B18,0)',MONEY),
("",None,None),
("Estimated hours",f'=IFERROR(INDEX({SVCH},MATCH(B4,{SVC},0)),0)+MAX(0,B6-2)*Settings!$E$18+MAX(0,B7-2)*Settings!$E$19+MAX(0,ROUNDUP((B5-5000)/1000,0))*Settings!$E$20+SUMIFS(Settings!$F$24:$F$29,Settings!$D$24:$D$29,"<>",E4:E9,"Yes")',"0.0"),
("Labor cost",'=B21*Settings!$B$8',MONEY),
("Profit before supplies & travel",'=B19-B22',MONEY),
("Effective rate per hour",'=IF(B21=0,0,B19/B21)',MONEY)]
q["A11"]="RESULT"; q["A11"].font=font(bold=True,color=ACC)
for i,(a,f,fm) in enumerate(out):
    r=12+i
    if not a: continue
    q.cell(r,1,a).font=font(bold=(a=="QUOTE PRICE"))
    c=q.cell(r,2,f); calc(c,fm,bold=(a=="QUOTE PRICE"))
q["A19"].font=font(bold=True,size=13,color=ACC); q["B19"].font=font(bold=True,size=13,color=ACC); q["B19"].fill=PatternFill("solid",fgColor=LIGHT)
q.cell(26,1,"Quote factors and add-on prices are set in Settings. The defaults are examples — set your own rates.").font=font(italic=True,color="57606A")

# ---------------- INVOICE
iv=sheet("Invoice")
title(iv,"Invoice","Type a Job ID in the yellow cell. Print or download as PDF.")
for col,w in zip("ABCDE",[22,30,14,14,16]): iv.column_dimensions[col].width=w
iv["D1"]="INVOICE"; iv["D1"].font=font(size=20,bold=True,color=ACC)
iv["A4"]="Job ID →"; iv["A4"].font=font(bold=True); inp(iv["B4"]); iv["B4"].value="J-0003"
M=f'MATCH($B$4,Jobs!$A$5:$A${4+N},0)'
def J(col): return f'=IFERROR(INDEX(Jobs!${col}$5:${col}${4+N},{M}),"")'
iv["A6"]="=Settings!B4"; iv["A6"].font=font(bold=True,size=13)
iv["A7"]="=Settings!B5"; iv["A8"]="=Settings!B6"; iv["A9"]="=Settings!B7"
iv["D6"]="Invoice #"; iv["E6"]='=IF(A12="","","INV-"&MID(B4,3,10))'
iv["D7"]="Service date"; iv["E7"]=J("B"); iv["E7"].number_format=DATE
iv["D8"]="Due date"; iv["E8"]='=IF(E7="","",E7+Settings!$B$11)'; iv["E8"].number_format=DATE
iv["D9"]="Status"; iv["E9"]=f'=IF(A12="","",IF({J("K")[1:]}="Yes","PAID","DUE"))'
iv["A11"]="BILL TO"; iv["A11"].font=font(bold=True,color=ACC)
iv["A12"]=f'=IFERROR(INDEX(Jobs!$C$5:$C${4+N},{M})&"","")'  # &"" so a blank cell shows blank, not 0; iv["A12"].font=font(bold=True)
iv["A13"]=f'=IFERROR(INDEX(Clients!$D$5:$D$304,MATCH(A12,Clients!$A$5:$A$304,0))&"","")'
iv["A14"]=f'=IFERROR(INDEX(Clients!$C$5:$C$304,MATCH(A12,Clients!$A$5:$A$304,0))&"","")'
for j,h in enumerate(["Description","","","","Amount"]):
    c=iv.cell(16,1+j,h or None); c.font=font(bold=True,color="FFFFFF"); c.fill=HDR_FILL
iv["A17"]=f'=IF(A12="","",IF({J("D")[1:]}<>"",{J("D")[1:]},IFERROR(IF(INDEX(Clients!$E$5:$E$304,MATCH(A12,Clients!$A$5:$A$304,0))&""="",NA(),INDEX(Clients!$E$5:$E$304,MATCH(A12,Clients!$A$5:$A$304,0))&""),"Lawn care service")))'
iv["E17"]=J("H"); iv["E17"].number_format=MONEY
iv["A18"]="Add-on services"; iv["E18"]=J("I"); iv["E18"].number_format=MONEY
iv["D20"]="TOTAL"; iv["D20"].font=font(bold=True,size=12); iv["E20"]='=IF(A12="","",N(E17)+N(E18))'; iv["E20"].number_format=MONEY; iv["E20"].font=font(bold=True,size=12)
iv["A23"]="Thank you for your business!"; iv["A23"].font=font(italic=True,color=ACC)
iv["A24"]=f'="Payment due within "&Settings!$B$11&" days."'
for r in range(6,25):
    for cc in range(1,6):
        cell=iv.cell(r,cc)
        if cell.font.name!=F or (cell.font.color is None): pass
        if not cell.font.bold and cell.font.size==11: cell.font=font()
iv.print_area="A1:E25"

# ---------------- DASHBOARD
d=sheet("Dashboard")
title(d,"Dashboard","Pick a year. Everything below updates from Jobs and Expenses.")
for col,w in zip("ABCDEFGHI",[14,13,13,13,13,13,9,13,13]): d.column_dimensions[col].width=w
d["A4"]="Year"; d["A4"].font=font(bold=True); inp(d["B4"],"0"); d["B4"].value=2026
# KPIs
kp=[("Revenue","=SUM(B9:B20)"),("Profit","=SUM(F9:F20)"),("Unpaid",f'=SUMIFS({JT},{JP},"No")'),("Tax set-aside","=MAX(0,B23)*Settings!$B$9")]
for i,(a,f) in enumerate(kp):
    col=2+i*2
    c=d.cell(5,col,a); c.font=font(bold=True,color="57606A",size=9)
    c=d.cell(6,col,f); c.font=font(bold=True,size=16,color=ACC); c.number_format=MONEY0
d.row_dimensions[6].height=28
header(d,8,["Month","Revenue","Job costs","Expenses","Jobs","Profit","Margin","Unpaid"])
for col,w in zip("ABCDEFGH",[14,13,13,13,9,13,9,13]): d.column_dimensions[col].width=w
for m in range(12):
    r=9+m; st_=f'DATE($B$4,{m+1},1)'; en=f'EOMONTH(DATE($B$4,{m+1},1),0)'
    calc(d.cell(r,1,f'=TEXT({st_},"mmm yyyy")'))
    calc(d.cell(r,2,f'=SUMIFS({JT},{JD},">="&{st_},{JD},"<="&{en})'),MONEY0)
    calc(d.cell(r,3,f'=SUMIFS(Jobs!$M$5:$M${4+N},{JD},">="&{st_},{JD},"<="&{en})+SUMIFS(Jobs!$O$5:$O${4+N},{JD},">="&{st_},{JD},"<="&{en})+SUMIFS(Jobs!$P$5:$P${4+N},{JD},">="&{st_},{JD},"<="&{en})'),MONEY0)
    calc(d.cell(r,4,f'=SUMIFS({XA},{XD},">="&{st_},{XD},"<="&{en})'),MONEY0)
    calc(d.cell(r,5,f'=COUNTIFS({JD},">="&{st_},{JD},"<="&{en})'),"0")
    calc(d.cell(r,6,f'=B{r}-C{r}-D{r}'),MONEY0)
    calc(d.cell(r,7,f'=IF(B{r}=0,0,F{r}/B{r})'),'0%;-0%;"-"')
    calc(d.cell(r,8,f'=SUMIFS({JT},{JD},">="&{st_},{JD},"<="&{en},{JP},"No")'),MONEY0)
    if m%2:
        for cc in range(1,9): d.cell(r,cc).fill=BAND
r=21; d.cell(r,1,"Total").font=font(bold=True)
for cc,col in zip(range(2,9),"BCDEFGH"):
    f = f'=IF(B21=0,0,F21/B21)' if col=="G" else f"=SUM({col}9:{col}20)"
    calc(d.cell(r,cc,f),'0%;-0%;"-"' if col=="G" else ("0" if col=="E" else MONEY0),bold=True)
d["A23"]="Profit for tax estimate"; d["A23"].font=font(); d["B23"]="=F21"; calc(d["B23"],MONEY0)
d["A24"]="Note: Tax set-aside is a planning estimate, not tax advice."; d["A24"].font=font(italic=True,size=9,color="57606A")
ch=BarChart(); ch.type="col"; ch.title="Revenue vs Profit by month"; ch.height=8; ch.width=20
ch.add_data(Reference(d,min_col=2,min_row=8,max_row=20),titles_from_data=True)
ch.add_data(Reference(d,min_col=6,min_row=8,max_row=20),titles_from_data=True)
ch.set_categories(Reference(d,min_col=1,min_row=9,max_row=20)); ch.y_axis.numFmt='$#,##0'
d.add_chart(ch,"J4")

# order tabs
order=["Start Here","Dashboard","Clients","Jobs","Expenses","Quote Calculator","Invoice","Settings"]
wb._sheets=[wb[n] for n in order]
wb.active=0
wb.save(sys.argv[1] if len(sys.argv)>1 else "Lawn-Care-Business-Manager.xlsx")
