"""Airbnb Host Business Manager: workbook generator.
Same style helpers as etsy-digital/_engine/build.py (the trade engine), but its own tabs:
Properties, Bookings, Turnovers, Expenses, Supplies, Monthly Summary, Pricing Calculator.
Run: python3 build.py [out.xlsx]"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from datetime import date

T = dict(name="Airbnb Host Business Manager", accent="1E6B64", light="E3F1EF")
ACC=T["accent"]; LIGHT=T["light"]
F="Arial"
def font(**k): k.setdefault("name",F); k.setdefault("size",10); return Font(**k)
HDR_FILL=PatternFill("solid",fgColor=ACC); IN_FILL=PatternFill("solid",fgColor="FFF7D6"); BAND=PatternFill("solid",fgColor=LIGHT)
HELP_FILL=PatternFill("solid",fgColor="F2F2F2")
thin=Side(style="thin",color="D0D7DE"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
BLUE="1F4EB4"; GREY="57606A"
MONEY='$#,##0.00;($#,##0.00);"-"'; MONEY0='$#,##0;($#,##0);"-"'; DATE='mmm d, yyyy'; PCT='0.0%'; PCT0='0%;-0%;"-"'
RED_FILL=PatternFill("solid",fgColor="F8D7DA"); RED_FONT=Font(name=F,color="9C1C28",bold=True); AMBER=PatternFill("solid",fgColor="FFF3CD")
N=1000      # booking / expense rows
NP=20       # property rows
P0,P1=5,4+NP          # Properties rows 5..24
B0,B1=5,4+N           # Bookings rows 5..1004

wb=Workbook()
def colw(ws,widths,start=1):
    from openpyxl.utils import get_column_letter
    for i,w in enumerate(widths,start): ws.column_dimensions[get_column_letter(i)].width=w
def sheet(t, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title=t; ws.sheet_view.showGridLines=False; ws.sheet_properties.tabColor=ACC
    return ws
def title(ws, text, sub):
    ws["A1"]=text; ws["A1"].font=font(size=18,bold=True,color=ACC)
    ws["A2"]=sub; ws["A2"].font=font(size=10,italic=True,color=GREY)
    ws.row_dimensions[1].height=30
def header(ws,row,cols,widths=None,start=1):
    for i,c in enumerate(cols,start):
        cell=ws.cell(row=row,column=i,value=c)
        cell.font=font(bold=True,color="FFFFFF"); cell.fill=HDR_FILL
        cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=BOX
    ws.row_dimensions[row].height=32
    if widths: colw(ws,widths,start)
def inp(c,fmt=None):
    c.font=font(color=BLUE); c.fill=IN_FILL; c.border=BOX
    if fmt: c.number_format=fmt
def calc(c,fmt=None,bold=False):
    c.font=font(bold=bold); c.border=BOX
    if fmt: c.number_format=fmt
def section(ws,ref,text):
    ws[ref]=text; ws[ref].font=font(bold=True,color=ACC)
def dv(ws,rng,formula,strict=True):
    v=DataValidation(type="list",formula1=formula,allow_blank=True)
    v.showErrorMessage=strict
    ws.add_data_validation(v); v.add(rng)

# ---------------- START HERE
s=sheet("Start Here",first=True)
title(s,T["name"],"Bookings · Turnovers · Occupancy & ADR · Expenses · Supplies · Profit dashboard: Google Sheets & Excel")
colw(s,[4,30,80])
steps=[("1. Settings","Your business name, tax set-aside %, length-of-stay discounts, booking channels with their fee %, expense categories and cleaner names."),
("2. Properties","One row per listing: nightly and weekend rate, cleaning fee, what you pay your cleaner, guests included and the extra-guest fee. Each row shows whether it's occupied today and the next check-in."),
("3. Bookings","Log every reservation: property, guest, channel, check-in and check-out. Nights, rent (weekday/weekend aware), cleaning fee, platform fee, payout and your net fill in. Type the real payout from your platform in 'Actual payout' to use it instead of the estimate. Double bookings are flagged."),
("4. Turnovers","Your cleaning schedule: every upcoming check-out, the next check-in at that property, the gap, and same-day turnover alerts. Assign the cleaner and tick Done in Bookings."),
("5. Expenses","Mortgage or rent, utilities, repairs, supplies, software, insurance. Pick the property, or leave it blank for costs shared across all properties."),
("6. Supplies","Par levels for toilet paper, coffee, soap and linens. Items at or below the reorder level turn red with the order quantity and cost."),
("7. Pricing Calculator","Price any stay before you accept it (rent, fees, platform cut, cleaning cost and net per night), and find the nightly rate you need to break even."),
("8. Dashboard & Monthly Summary","Pick a year. See payout, profit, occupancy, ADR and RevPAR per property and by month. Monthly Summary drills into one property.")]
section(s,"B4","HOW TO USE"); s["B4"].font=font(bold=True,size=12,color=ACC)
for i,(a,b) in enumerate(steps):
    r=5+i; s.cell(r,2,a).font=font(bold=True); c=s.cell(r,3,b); c.font=font(); c.alignment=Alignment(wrap_text=True,vertical="top"); s.row_dimensions[r].height=42
r=5+len(steps)+1
s.cell(r,2,"COLOR KEY").font=font(bold=True,size=12,color=ACC)
k=s.cell(r+1,2,"Yellow cell, blue text"); inp(k); s.cell(r+1,3,"Type here: your inputs").font=font()
k=s.cell(r+2,2,"White cell, black text"); calc(k); s.cell(r+2,3,"Calculated automatically. Don't type over these.").font=font()
s.cell(r+4,2,"KEY TERMS").font=font(bold=True,size=12,color=ACC)
terms=[("Occupancy","Booked nights ÷ nights available."),("ADR","Average daily rate: rent ÷ booked nights (cleaning and other fees not included)."),
("RevPAR","Rent ÷ nights available. Combines price and occupancy in one number."),("Payout","What the platform pays you after its host fee (or the actual payout you type in)."),
("Split months","A stay that crosses into a new month is split by nights, so each month gets its fair share.")]
for i,(a,b) in enumerate(terms):
    s.cell(r+5+i,2,a).font=font(bold=True); s.cell(r+5+i,3,b).font=font()
r=r+5+len(terms)+1
s.cell(r,2,"GOOGLE SHEETS").font=font(bold=True,size=12,color=ACC)
c=s.cell(r+1,3,"Upload this file to Google Drive, open it, then File > Save as Google Sheets. Everything works in both Excel and Google Sheets."); c.font=font(); c.alignment=Alignment(wrap_text=True); s.row_dimensions[r+1].height=30
c=s.cell(r+2,3,"The example rows (marked EXAMPLE) show the format. Delete them when you start. After that, clear a booking row instead of deleting it, so Booking IDs stay the same.")
c.font=font(italic=True,color=GREY); c.alignment=Alignment(wrap_text=True); s.row_dimensions[r+2].height=30
c=s.cell(r+3,3,"Default fees and rates are examples. Check your own platform fees. Expense categories are an organizing aid, not tax advice. Not affiliated with or endorsed by Airbnb, Inc. or Vrbo.")
c.font=font(italic=True,color=GREY); c.alignment=Alignment(wrap_text=True); s.row_dimensions[r+3].height=30

# ---------------- SETTINGS
st=sheet("Settings")
title(st,"Settings","Your business details and defaults. Every yellow cell is editable.")
colw(st,[36,18,4,26,14,4,4,26,4,22,4,22])
rows=[("Business name","Harbor Stays",None),("Owner / contact","Your Name",None),("Phone","(555) 555-0123",None),("Email","hello@example.com",None),
("Tax set-aside % of profit",0.25,PCT),("Weekly discount (7+ nights)",0.10,PCT),("Monthly discount (28+ nights)",0.20,PCT),("Check-out time","11:00 AM",None),("Check-in time","4:00 PM",None)]
for i,(a,b,f) in enumerate(rows):
    r=4+i; st.cell(r,1,a).font=font(); inp(st.cell(r,2,b),f)
st["B8"].comment=Comment("A planning estimate only. Ask a tax professional what to set aside.","Template")
st["B9"].comment=Comment("Applied to rent when a stay is 7-27 nights and you haven't typed a nightly rate override. Set 0% if you don't offer it.","Template")
st["B10"].comment=Comment("Applied to rent when a stay is 28+ nights. Set 0% if you don't offer it.","Template")
section(st,"D3","BOOKING CHANNELS")
for j,h in enumerate(["Channel","Host fee %"]):
    c=st.cell(4,4+j,h); c.font=font(bold=True,color="FFFFFF"); c.fill=HDR_FILL; c.border=BOX
chans=[("Airbnb",0.03),("Airbnb (host-only fee)",0.155),("Vrbo",0.08),("Booking.com",0.15),("Direct booking",0.03),("",None),("",None),("",None)]
for i,(a,b) in enumerate(chans):
    r=5+i; inp(st.cell(r,4,a or None)); inp(st.cell(r,5,b),PCT)
st["E5"].comment=Comment("Example: Airbnb split-fee hosts pay about 3%. Check the fee shown on your own payouts and change it here.","Template")
st["E6"].comment=Comment("Example: host-only fee is about 14-16%. Check your own payouts.","Template")
st["E7"].comment=Comment("Example: 5% commission + 3% payment processing on pay-per-booking. Check your own plan.","Template")
st["E9"].comment=Comment("Example: card processing fee for bookings you take yourself.","Template")
section(st,"H3","EXPENSE CATEGORIES")
cats=["Mortgage interest / rent","Utilities","Internet & streaming","Cleaning & laundry","Supplies & consumables","Repairs & maintenance","Furnishings & equipment",
"Insurance","Property tax","HOA fees","Software & pricing tools","Photos & advertising","Lodging taxes paid","Other",""]
for i,a in enumerate(cats): inp(st.cell(4+i,8,a or None))
st.cell(20,8,"Categories are an organizing aid, not tax advice.").font=font(italic=True,size=8,color=GREY)
section(st,"J3","CLEANERS")
for i,a in enumerate(["Me","Maria (cleaner)","Sparkle Turnover Co.","","","",""]): inp(st.cell(4+i,10,a or None))
section(st,"J13","PAYMENT METHODS")
for i,a in enumerate(["Platform payout","Card","Bank transfer","Zelle","Venmo","Cash","Check"]): inp(st.cell(14+i,10,a or None))
section(st,"L3","PROPERTY PICKER (auto)")
st["L4"]="All properties"; calc(st["L4"])
for i in range(NP):
    calc(st.cell(5+i,12,f'=IF(Properties!A{P0+i}="","",Properties!A{P0+i})'))
st.cell(26,12,"Feeds the property dropdowns. Don't edit.").font=font(italic=True,size=8,color=GREY)

CH="Settings!$D$5:$D$12"; CHF="Settings!$E$5:$E$12"
PN=f"Properties!$A${P0}:$A${P1}"
def PR(col): return f"Properties!${col}${P0}:${col}${P1}"
def BR(col): return f"Bookings!${col}${B0}:${col}${B1}"
def XR(col): return f"Expenses!${col}$5:${col}${4+N}"

# ---------------- PROPERTIES
pp=sheet("Properties")
title(pp,"Properties","One row per listing. Rates and fees here feed Bookings and the Pricing Calculator.")
cols=["Property name","Address","Bedrooms","Max guests","Nightly rate (Sun-Thu)","Weekend rate (Fri-Sat)","Cleaning fee (guest pays)","Cleaner pay per turnover",
"Guests included","Extra guest fee / night","Listed since","Wi-Fi, lockbox, notes","Today","Next check-in","Upcoming stays"]
header(pp,4,cols,[24,26,10,9,12,12,12,12,10,11,13,30,13,13,10])
pex=[("EXAMPLE – Lakeview Cabin","12 Shore Rd",2,6,165,210,95,65,4,15,date(2025,6,1),"Lockbox 4471 · Wi-Fi LakeView5G")]
for r in range(P0,P1+1):
    e=pex[r-P0] if r-P0<len(pex) else [None]*12
    for j in range(12):
        fm = MONEY if j in (4,5,6,7,9) else DATE if j==10 else "0" if j in (2,3,8) else None
        inp(pp.cell(r,1+j,e[j]),fm)
    A=f"$A{r}"
    calc(pp.cell(r,13,f'=IF({A}="","",IF(COUNTIFS({BR("D")},{A},{BR("B")},"<="&TODAY(),{BR("C")},">"&TODAY(),{BR("AA")},1)>0,"Occupied",IF(COUNTIFS({BR("D")},{A},{BR("C")},TODAY(),{BR("AA")},1)>0,"Turnover today","Vacant")))'))
    calc(pp.cell(r,14,f'=IF({A}="","",IF(_xlfn.MINIFS({BR("B")},{BR("D")},{A},{BR("B")},">="&TODAY(),{BR("AA")},1)=0,"None",_xlfn.MINIFS({BR("B")},{BR("D")},{A},{BR("B")},">="&TODAY(),{BR("AA")},1)))'),DATE)
    calc(pp.cell(r,15,f'=IF({A}="","",COUNTIFS({BR("D")},{A},{BR("B")},">="&TODAY(),{BR("AA")},1))'),"0")
pp["K4"].comment=Comment("First night this property was available to book. Occupancy counts available nights from this date. Leave blank if it was listed before the year you're looking at.","Template")
pp.conditional_formatting.add(f"M{P0}:M{P1}",CellIsRule(operator="equal",formula=['"Turnover today"'],fill=AMBER,font=Font(name=F,bold=True)))
pp.conditional_formatting.add(f"M{P0}:M{P1}",CellIsRule(operator="equal",formula=['"Occupied"'],fill=BAND,font=Font(name=F,color=ACC,bold=True)))
pp.freeze_panes="B5"

# ---------------- BOOKINGS
bk=sheet("Bookings")
title(bk,"Bookings","Log every reservation. Yellow = you type, white = calculated. Type the real payout in 'Actual payout' when you have it.")
cols=["Booking ID","Check-in","Check-out","Property","Guest name","Channel","Guests","Nights","Nightly rate override","Rent","Cleaning fee","Extra guest fees",
"Other fees / discounts","Guest total","Platform fee","Actual payout (optional)","Payout","Cleaning cost","Net after cleaning","Status","Paid?","Cleaner","Clean done?","Notes","Check",
"Turnover key","Include","Rent / night","Payout / night","Cleaning cost / night","Platform fee / night","Cleaning fee / night","Check-in key","Nights (helper)"]
header(bk,4,cols,[10,12,12,22,18,15,7,7,10,11,10,10,11,11,10,11,11,10,11,10,7,15,8,24,13,9,7,9,9,9,9,9,14,8])
bex=[(date(2026,10,2),date(2026,10,5),"EXAMPLE – Lakeview Cabin","Jordan P.","Airbnb",4,None,None,None,"Booked","Yes","Maria (cleaner)","Yes",None),
(date(2026,10,9),date(2026,10,11),"EXAMPLE – Lakeview Cabin","Sam & Lee R.","Direct booking",6,None,25,None,"Booked","No","Maria (cleaner)",None,"Pet fee $25")]
for i in range(N):
    r=B0+i; e=bex[i] if i<len(bex) else [None]*14
    PM=f'MATCH(D{r},{PN},0)'
    def P(col): return f'INDEX({PR(col)},{PM})'
    calc(bk.cell(r,1,f'=IF(B{r}="","","B-"&TEXT(ROW()-4,"0000"))'))
    inp(bk.cell(r,2,e[0]),DATE); inp(bk.cell(r,3,e[1]),DATE); inp(bk.cell(r,4,e[2])); inp(bk.cell(r,5,e[3])); inp(bk.cell(r,6,e[4])); inp(bk.cell(r,7,e[5]),"0")
    calc(bk.cell(r,8,f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r})),MAX(0,C{r}-B{r}),"")'),"0")
    inp(bk.cell(r,9,e[6]),MONEY)
    wk=f'(INT((WEEKDAY(B{r}-6)+H{r}-1)/7)+INT((WEEKDAY(B{r}-7)+H{r}-1)/7))'
    disc=f'IF(H{r}>=28,Settings!$B$10,IF(H{r}>=7,Settings!$B$9,0))'
    calc(bk.cell(r,10,f'=IF(OR(D{r}="",N(H{r})=0),"",IFERROR(IF(I{r}<>"",H{r}*I{r},ROUND(((H{r}-{wk})*{P("E")}+{wk}*{P("F")})*(1-{disc}),2)),0))'),MONEY)
    calc(bk.cell(r,11,f'=IF(J{r}="","",IFERROR(N({P("G")}),0))'),MONEY)
    calc(bk.cell(r,12,f'=IF(J{r}="","",IFERROR(MAX(0,N(G{r})-{P("I")})*N({P("J")})*H{r},0))'),MONEY)
    inp(bk.cell(r,13,e[7]),MONEY)
    calc(bk.cell(r,14,f'=IF(J{r}="","",J{r}+K{r}+L{r}+N(M{r}))'),MONEY)
    calc(bk.cell(r,15,f'=IF(N{r}="","",IF(P{r}<>"",N{r}-P{r},ROUND(N{r}*IFERROR(INDEX({CHF},MATCH(F{r},{CH},0)),0),2)))'),MONEY)
    inp(bk.cell(r,16,e[8]),MONEY)
    calc(bk.cell(r,17,f'=IF(N{r}="","",IF(P{r}<>"",P{r},N{r}-O{r}))'),MONEY)
    calc(bk.cell(r,18,f'=IF(J{r}="","",IFERROR(N({P("H")}),0))'),MONEY)
    calc(bk.cell(r,19,f'=IF(Q{r}="","",Q{r}-R{r})'),MONEY)
    inp(bk.cell(r,20,e[9])); inp(bk.cell(r,21,e[10])); inp(bk.cell(r,22,e[11])); inp(bk.cell(r,23,e[12])); inp(bk.cell(r,24,e[13]))
    calc(bk.cell(r,25,f'=IF(OR(D{r}="",B{r}="",C{r}=""),"",IF(N(C{r})<=N(B{r}),"Check dates",IF(T{r}="Cancelled","",IF(COUNTIFS($D${B0}:$D${B1},D{r},$B${B0}:$B${B1},"<"&C{r},$C${B0}:$C${B1},">"&B{r},$T${B0}:$T${B1},"<>Cancelled")>1,"Double-booked",""))))'))
    # helpers (grey)
    for cc,f,fm in [(26,f'=IF(AND(AA{r}=1,ISNUMBER(C{r})),C{r}+ROW()/100000,"")',"0.00000"),
                    (27,f'=IF(AND(D{r}<>"",ISNUMBER(B{r}),ISNUMBER(C{r}),N(H{r})>0,T{r}<>"Cancelled",J{r}<>""),1,0)',"0"),
                    (28,f'=IF(AA{r}=1,J{r}/H{r},0)',MONEY),(29,f'=IF(AA{r}=1,Q{r}/H{r},0)',MONEY),(30,f'=IF(AA{r}=1,R{r}/H{r},0)',MONEY),
                    (31,f'=IF(AA{r}=1,O{r}/H{r},0)',MONEY),(32,f'=IF(AA{r}=1,K{r}/H{r},0)',MONEY),
                    (33,f'=IF(AA{r}=1,D{r}&"|"&TEXT(B{r},"0"),"")',None),(34,f'=IF(AA{r}=1,H{r},0)',"0")]:
        c=bk.cell(r,cc,f); c.font=font(color="8C959F",size=9); c.fill=HELP_FILL
        if fm: c.number_format=fm
    if i%2:
        for cc in (1,8,10,11,12,14,15,17,18,19): bk.cell(r,cc).fill=BAND
bk.cell(3,26,"Helper columns: used by the dashboard and turnovers. Don't edit.").font=font(italic=True,size=8,color=GREY)
dv(bk,f"D{B0}:D{B1}",PN); dv(bk,f"F{B0}:F{B1}",CH); dv(bk,f"T{B0}:T{B1}",'"Booked,Cancelled"'); dv(bk,f"U{B0}:U{B1}",'"Yes,No"')
dv(bk,f"V{B0}:V{B1}","Settings!$J$4:$J$10"); dv(bk,f"W{B0}:W{B1}",'"Yes,No"')
bk.conditional_formatting.add(f"Y{B0}:Y{B1}",CellIsRule(operator="notEqual",formula=['""'],fill=RED_FILL,font=RED_FONT))
bk.conditional_formatting.add(f"U{B0}:U{B1}",CellIsRule(operator="equal",formula=['"No"'],fill=RED_FILL,font=RED_FONT))
bk.conditional_formatting.add(f"T{B0}:T{B1}",CellIsRule(operator="equal",formula=['"Cancelled"'],font=Font(name=F,color="8C959F",strike=True)))
bk.freeze_panes="E5"; bk.auto_filter.ref=f"A4:Y{B1}"
bk["I4"].comment=Comment("Leave blank to use the property's weekday/weekend rates (with your length-of-stay discount). Type a rate to use one flat nightly rate for this stay.","Template")
bk["M4"].comment=Comment("Pet fee, early check-in, damage charge (+) or a discount you gave (-).","Template")
bk["P4"].comment=Comment("Optional. Type the payout your platform actually sent. It replaces the estimate, and the platform fee becomes the difference.","Template")
bk["U4"].comment=Comment("For direct bookings: No = the guest still owes you. Unpaid totals show on the Dashboard.","Template")

# ---------------- TURNOVERS
tv=sheet("Turnovers")
title(tv,"Turnovers","Your cleaning schedule: the next 40 check-outs, sorted by date. Assign cleaners and tick Done in Bookings.")
colw(tv,[13,7,22,18,13,18,8,10,20,17,8,30])
tv["A4"]="Show from"; tv["A4"].font=font(bold=True); inp(tv["B4"],DATE); tv["B4"].value="=TODAY()"
tv.merge_cells("B4:C4")
tv["D4"]=f'="Check-out "&Settings!$B$11&" · Check-in "&Settings!$B$12'; tv["D4"].font=font(italic=True,color=GREY)
tv.cell(5,1,"Leave 'Show from' as =TODAY(), or type a date to look ahead or back.").font=font(italic=True,size=8,color=GREY)
header(tv,6,["Turnover date","Day","Property","Checking out","Next check-in","Next guest","Guests","Gap (days)","Alert","Cleaner","Done?","Booking notes"])
for k in range(40):
    r=7+k
    key=f'SMALL({BR("Z")},COUNTIF({BR("Z")},"<"&$B$4)+{k+1})'
    calc(tv.cell(r,13,f'=IFERROR(MATCH({key},{BR("Z")},0),"")'),"0"); tv.cell(r,13).font=font(color="8C959F",size=8)
    def IX(col): return f'INDEX({BR(col)},$M{r})'
    calc(tv.cell(r,1,f'=IF($M{r}="","",{IX("C")})'),DATE)
    calc(tv.cell(r,2,f'=IF(A{r}="","",TEXT(A{r},"ddd"))'))
    calc(tv.cell(r,3,f'=IF($M{r}="","",{IX("D")})'))
    calc(tv.cell(r,4,f'=IF($M{r}="","",{IX("E")}&"")'))
    calc(tv.cell(r,5,f'=IF($M{r}="","",IF(_xlfn.MINIFS({BR("B")},{BR("D")},C{r},{BR("B")},">="&A{r},{BR("AA")},1)=0,"",_xlfn.MINIFS({BR("B")},{BR("D")},C{r},{BR("B")},">="&A{r},{BR("AA")},1)))'),DATE)
    nm=f'MATCH(C{r}&"|"&TEXT(E{r},"0"),{BR("AG")},0)'
    calc(tv.cell(r,6,f'=IF(E{r}="","",IFERROR(INDEX({BR("E")},{nm})&"",""))'))
    calc(tv.cell(r,7,f'=IF(E{r}="","",IFERROR(INDEX({BR("G")},{nm}),""))'),"0")
    calc(tv.cell(r,8,f'=IF(OR(A{r}="",E{r}=""),"",E{r}-A{r})'),"0")
    calc(tv.cell(r,9,f'=IF(A{r}="","",IF(E{r}="","No next booking",IF(H{r}=0,"SAME-DAY turnover",IF(H{r}=1,"Next-day check-in",""))))'))
    calc(tv.cell(r,10,f'=IF($M{r}="","",IF({IX("V")}="","Unassigned",{IX("V")}))'))
    calc(tv.cell(r,11,f'=IF($M{r}="","",IF({IX("W")}="Yes","Yes",""))'))
    calc(tv.cell(r,12,f'=IF($M{r}="","",{IX("X")}&"")'))
    if k%2:
        for cc in range(1,13): tv.cell(r,cc).fill=BAND
tv.cell(6,13,"idx").font=font(color="8C959F",size=8)
tv.column_dimensions["M"].width=5
tv.conditional_formatting.add("I7:I46",CellIsRule(operator="equal",formula=['"SAME-DAY turnover"'],fill=RED_FILL,font=RED_FONT))
tv.conditional_formatting.add("I7:I46",CellIsRule(operator="equal",formula=['"Next-day check-in"'],fill=AMBER))
tv.conditional_formatting.add("J7:J46",CellIsRule(operator="equal",formula=['"Unassigned"'],font=Font(name=F,color="9C1C28",bold=True)))
tv.conditional_formatting.add("K7:K46",CellIsRule(operator="equal",formula=['"Yes"'],font=Font(name=F,color=ACC,bold=True)))
tv.freeze_panes="A7"; tv.print_area="A1:L46"

# ---------------- EXPENSES
xp=sheet("Expenses")
title(xp,"Expenses","Every cost of running your rentals. Pick the property, or leave it blank for costs shared by all.")
header(xp,4,["Date","Property (blank = shared)","Category","Vendor","Description","Amount","Paid with","Receipt saved?"],[13,24,24,20,30,12,15,12])
xex=[(date(2026,10,1),"EXAMPLE – Lakeview Cabin","Utilities","City Power","EXAMPLE – Electric bill",142.18,"Card","Yes"),
(date(2026,10,4),None,"Software & pricing tools","PriceLabs","EXAMPLE – Dynamic pricing (all listings)",19.99,"Card","Yes")]
for i in range(N):
    r=5+i; e=xex[i] if i<len(xex) else [None]*8
    for j in range(8): inp(xp.cell(r,1+j,e[j]),DATE if j==0 else MONEY if j==5 else None)
dv(xp,f"B5:B{4+N}",PN); dv(xp,f"C5:C{4+N}","Settings!$H$4:$H$18"); dv(xp,f"G5:G{4+N}","Settings!$J$14:$J$20"); dv(xp,f"H5:H{4+N}",'"Yes,No"')
xp.freeze_panes="A5"; xp.auto_filter.ref=f"A4:H{4+N}"

# ---------------- SUPPLIES
sp=sheet("Supplies")
title(sp,"Supplies","Set a par level for each item. Count what's on hand after a turnover; anything at or below its reorder level turns red.")
colw(sp,[26,22,10,9,9,10,11,10,10,11,20,13])
sp["A4"]="Reorder total"; sp["A4"].font=font(bold=True)
calc(sp["B4"],MONEY,bold=True); sp["B4"].value="=SUM(J7:J106)"; sp["B4"].font=font(bold=True,size=12,color=ACC)
sp["D4"]='=COUNTIF(G7:G106,"Reorder")&" item(s) to reorder"'; sp["D4"].font=font(bold=True,color="9C1C28")
header(sp,6,["Item","Property (blank = all)","Unit","Par level","On hand","Reorder at","Status","Order qty","Unit cost","Order cost","Where to buy","Last restocked"])
sup=[("Toilet paper","roll",24,30,12,0.55,"Costco"),("Paper towels","roll",8,10,4,1.10,"Costco"),("Coffee pods","pod",40,18,20,0.45,"Amazon"),
("Dish soap","bottle",2,2,1,3.50,"Target"),("Dishwasher pods","pod",30,24,10,0.30,"Costco"),("Laundry detergent","pod",40,35,15,0.28,"Costco"),
("Trash bags (13 gal)","bag",40,16,15,0.20,"Costco"),("Shampoo & conditioner minis","set",12,9,6,1.25,"Amazon"),("Hand soap refill","bottle",3,3,1,4.25,"Target"),
("Sponges","pack",2,1,1,4.00,"Target"),("Spare light bulbs","bulb",4,4,2,2.50,"Home Depot"),("AA batteries (remotes)","battery",8,4,4,0.60,"Amazon")]
for i in range(100):
    r=7+i; e=sup[i] if i<len(sup) else None
    vals=[e[0],None,e[1],e[2],e[3],e[4],None,None,e[5],None,e[6],None] if e else [None]*12
    for j in (0,1,2,3,4,5,8,10,11):
        inp(sp.cell(r,1+j,vals[j]),"0" if j in (3,4,5) else MONEY if j==8 else DATE if j==11 else None)
    calc(sp.cell(r,7,f'=IF(OR(A{r}="",E{r}=""),"",IF(E{r}<=F{r},"Reorder","OK"))'))
    calc(sp.cell(r,8,f'=IF(G{r}="Reorder",MAX(0,D{r}-E{r}),"")'),"0")
    calc(sp.cell(r,10,f'=IF(H{r}="","",H{r}*N(I{r}))'),MONEY)
dv(sp,"B7:B106",PN)
sp.conditional_formatting.add("G7:G106",CellIsRule(operator="equal",formula=['"Reorder"'],fill=RED_FILL,font=RED_FONT))
sp.conditional_formatting.add("G7:G106",CellIsRule(operator="equal",formula=['"OK"'],font=Font(name=F,color=ACC,bold=True)))
sp.freeze_panes="B7"
sp.cell(108,1,"Starter list: edit items, quantities and prices to match your properties.").font=font(italic=True,color=GREY)

# ---------------- PRICING CALCULATOR
q=sheet("Pricing Calculator")
title(q,"Pricing Calculator","Price a stay before you accept it, and find the nightly rate you need to break even.")
colw(q,[34,18,4,38,18])
section(q,"A4","PRICE A STAY")
qi=[("Property","=Properties!A5",None),("Check-in date",date(2026,11,13),DATE),("Nights",3,"0"),("Guests",5,"0"),("Channel","Airbnb",None),("Other fees / discounts",0,MONEY)]
for i,(a,b,f) in enumerate(qi):
    r=5+i; q.cell(r,1,a).font=font(); inp(q.cell(r,2,b),f)
q["B5"].comment=Comment("Pick a property from the dropdown.","Template")
dv(q,"B5",PN); dv(q,"B9",CH)
def QP(col): return f'IFERROR(INDEX({PR(col)},MATCH($B$5,{PN},0)),0)'
wkq='(INT((WEEKDAY(B6-6)+B7-1)/7)+INT((WEEKDAY(B6-7)+B7-1)/7))'
out=[("Weekend nights (Fri, Sat)",f'=IF(B7<=0,0,{wkq})',"0"),
("Weekday nights",'=MAX(0,B7-B12)',"0"),
("Rent before discount",f'=B13*{QP("E")}+B12*{QP("F")}',MONEY),
("Length-of-stay discount",'=-B14*IF(B7>=28,Settings!$B$10,IF(B7>=7,Settings!$B$9,0))',MONEY),
("Rent",'=B14+B15',MONEY),
("Cleaning fee",f'={QP("G")}',MONEY),
("Extra guest fees",f'=MAX(0,B8-{QP("I")})*{QP("J")}*B7',MONEY),
("Other fees / discounts",'=B10',MONEY),
("GUEST PAYS (before taxes)",'=SUM(B16:B19)',MONEY),
("Platform fee",f'=-ROUND(B20*IFERROR(INDEX({CHF},MATCH(B9,{CH},0)),0),2)',MONEY),
("Your payout",'=B20+B21',MONEY),
("Cleaner pay",f'=-{QP("H")}',MONEY),
("NET TO YOU",'=B22+B23',MONEY),
("Net per night",'=IF(B7=0,0,B24/B7)',MONEY)]
for i,(a,f,fm) in enumerate(out):
    r=12+i; q.cell(r,1,a).font=font(bold=a.isupper()); calc(q.cell(r,2,f),fm,bold=a.isupper())
q["A11"]="RESULT"; q["A11"].font=font(bold=True,color=ACC)
for ref in ("A24","B24"): q[ref].font=font(bold=True,size=13,color=ACC)
q["B24"].fill=PatternFill("solid",fgColor=LIGHT)
section(q,"D4","BREAK-EVEN NIGHTLY RATE")
bi=[("Monthly fixed costs (mortgage, utilities, internet, insurance…)",2400,MONEY),("Expected occupancy",0.65,PCT0),("Average stay (nights)",3,"0.0"),("Target monthly profit",1000,MONEY),("Channel","Airbnb",None)]
for i,(a,b,f) in enumerate(bi):
    r=5+i; c=q.cell(r,4,a); c.font=font(); c.alignment=Alignment(wrap_text=True); inp(q.cell(r,5,b),f)
q.row_dimensions[5].height=26
dv(q,"E9",CH)
bo=[("Booked nights per month",'=30.4*E6',"0.0"),("Stays (turnovers) per month",'=IF(E7=0,0,D12_)',"0.0"),
("Platform fee %",f'=IFERROR(INDEX({CHF},MATCH(E9,{CH},0)),0)',PCT),
("BREAK-EVEN RATE / night",'=IF(E12=0,0,((E5+E13*CLEAN)/(1-E14)-E13*CFEE)/E12)',MONEY),
("RATE FOR TARGET PROFIT",'=IF(E12=0,0,((E5+E8+E13*CLEAN)/(1-E14)-E13*CFEE)/E12)',MONEY),
("Uses the property picked on the left for cleaning fee and cleaner pay.",None,None)]
q["D11"]="RESULT"; q["D11"].font=font(bold=True,color=ACC)
for i,(a,f,fm) in enumerate(bo):
    r=12+i; c=q.cell(r,4,a)
    if f is None: c.font=font(italic=True,size=9,color=GREY); continue
    c.font=font(bold=a.isupper())
    f=f.replace("D12_","E12/E7").replace("CLEAN",QP("H")).replace("CFEE",QP("G"))
    calc(q.cell(r,5,f),fm,bold=a.isupper())
for ref in ("D15","E15","D16","E16"): q[ref].font=font(bold=True,size=12,color=ACC)
q["E15"].fill=PatternFill("solid",fgColor=LIGHT); q["E16"].fill=PatternFill("solid",fgColor=LIGHT)
q.cell(27,1,"Rates, fees and discounts come from Properties and Settings. Defaults are examples: set your own. Lodging taxes are not included.").font=font(italic=True,color=GREY)

# ---------------- monthly metric formulas (shared by Dashboard and Monthly Summary)
def nim(ms,me1):
    """nights of each booking inside [ms, me1): array expression for SUMPRODUCT"""
    b,c=BR("B"),BR("C")
    s=f'(({b}+{ms}+ABS({b}-{ms}))/2)'; e=f'(({c}+{me1}-ABS({c}-{me1}))/2)'
    return f'((({e}-{s})+ABS({e}-{s}))/2)'
def avail(ms,me1,pfilter):
    g=PR("K"); s=f'(({g}+{ms}+ABS({g}-{ms}))/2)'
    return f'SUMPRODUCT(({PN}<>"")*{pfilter}*((({me1}-{s})+ABS({me1}-{s}))/2))'
def metrics(ms,me1,sel):
    """sel: cell ref holding a property name or 'All properties'"""
    bf=f'(({sel}="All properties")+({BR("D")}={sel})>0)'
    pf=f'(({sel}="All properties")+({PN}={sel})>0)'
    base=f'{BR("AA")}*{bf}*{nim(ms,me1)}'
    return dict(avail=avail(ms,me1,pf),
        nights=f'SUMPRODUCT({base})',
        rent=f'SUMPRODUCT({base}*{BR("AB")})', payout=f'SUMPRODUCT({base}*{BR("AC")})', clean=f'SUMPRODUCT({base}*{BR("AD")})',
        fees=f'SUMPRODUCT({base}*{BR("AE")})', cfee=f'SUMPRODUCT({base}*{BR("AF")})',
        exp=f'IF({sel}="All properties",SUMIFS({XR("F")},{XR("A")},">="&{ms},{XR("A")},"<"&{me1}),SUMIFS({XR("F")},{XR("A")},">="&{ms},{XR("A")},"<"&{me1},{XR("B")},{sel}))',
        stays=f'SUMPRODUCT({BR("AA")}*{bf}*({BR("B")}>={ms})*({BR("B")}<{me1}))',
        staynights=f'SUMPRODUCT({BR("AA")}*{bf}*({BR("B")}>={ms})*({BR("B")}<{me1})*{BR("AH")})')

# ---------------- DASHBOARD
d=sheet("Dashboard")
title(d,"Dashboard","Pick a year. Everything updates from Properties, Bookings and Expenses.")
colw(d,[24,11,11,11,11,13,13,13,13,4])
d["A4"]="Year"; d["A4"].font=font(bold=True); inp(d["B4"],"0"); d["B4"].value=2026
d["D4"]="All properties"; d["D4"].font=font(color="FFFFFF"); d["D4"].number_format=';;;'   # filter value used by formulas
kp=[("Payout","=F22"),("Profit","=I22"),("Occupancy","=C22"),("ADR (avg nightly rate)","=D22"),("Unpaid direct bookings",f'=SUMIFS({BR("Q")},{BR("U")},"No",{BR("AA")},1)'),("Tax set-aside","=MAX(0,I22)*Settings!$B$8")]
tiles=[("A",kp[0]),("C",kp[1]),("E",kp[2]),("F",kp[3]),("H",kp[4]),("I",kp[5])]
for colL,(a,f) in tiles:
    c=d[f"{colL}6"]; c.value=a; c.font=font(bold=True,color=GREY,size=9)
    c=d[f"{colL}7"]; c.value=f; c.font=font(bold=True,size=16,color=ACC); c.number_format=PCT0 if a=="Occupancy" else MONEY0
d.row_dimensions[7].height=28
header(d,9,["Month","Nights booked","Occupancy","ADR","RevPAR","Payout","Cleaning costs","Expenses","Profit"])
YS='DATE($B$4,1,1)'; YE='DATE($B$4+1,1,1)'
def month_rows(ws,r0,yref,sel,extra=False):
    for m in range(12):
        r=r0+m; ms=f'DATE({yref},{m+1},1)'; me=f'DATE({yref},{m+2},1)'
        M=metrics(ms,me,sel)
        calc(ws.cell(r,1,f'=TEXT({ms},"mmm yyyy")'))
        calc(ws.cell(r,2,'='+M["nights"]),"0")
        ws.cell(r,11,'='+M["avail"]); ws.cell(r,11).font=font(color="8C959F",size=8); ws.cell(r,11).number_format="0"
        calc(ws.cell(r,3,f'=IF(K{r}=0,0,B{r}/K{r})'),PCT0)
        ws.cell(r,12,'='+M["rent"]); ws.cell(r,12).font=font(color="8C959F",size=8); ws.cell(r,12).number_format=MONEY0
        calc(ws.cell(r,4,f'=IF(B{r}=0,0,L{r}/B{r})'),MONEY0)
        calc(ws.cell(r,5,f'=IF(K{r}=0,0,L{r}/K{r})'),MONEY0)
        calc(ws.cell(r,6,'='+M["payout"]),MONEY0)
        calc(ws.cell(r,7,'='+M["clean"]),MONEY0)
        calc(ws.cell(r,8,'='+M["exp"]),MONEY0)
        calc(ws.cell(r,9,f'=F{r}-G{r}-H{r}'),MONEY0)
        if extra:
            calc(ws.cell(r,13,'='+M["stays"]),"0")
            calc(ws.cell(r,14,f'=IF(M{r}=0,0,{M["staynights"]}/M{r})'),"0.0")
            calc(ws.cell(r,15,'='+M["cfee"]),MONEY0)
            calc(ws.cell(r,16,'='+M["fees"]),MONEY0)
        if m%2:
            for cc in range(1,10 if not extra else 17):
                if cc in (11,12): continue
                ws.cell(r,cc).fill=BAND
    r=r0+12; ws.cell(r,1,"Year total").font=font(bold=True)
    calc(ws.cell(r,2,f'=SUM(B{r0}:B{r-1})'),"0",True)
    ws.cell(r,11,f'=SUM(K{r0}:K{r-1})'); ws.cell(r,11).font=font(color="8C959F",size=8); ws.cell(r,11).number_format="0"
    ws.cell(r,12,f'=SUM(L{r0}:L{r-1})'); ws.cell(r,12).font=font(color="8C959F",size=8); ws.cell(r,12).number_format=MONEY0
    calc(ws.cell(r,3,f'=IF(K{r}=0,0,B{r}/K{r})'),PCT0,True)
    calc(ws.cell(r,4,f'=IF(B{r}=0,0,L{r}/B{r})'),MONEY0,True)
    calc(ws.cell(r,5,f'=IF(K{r}=0,0,L{r}/K{r})'),MONEY0,True)
    for col in "FGHI": calc(ws[f"{col}{r}"],MONEY0,True); ws[f"{col}{r}"].value=f"=SUM({col}{r0}:{col}{r-1})"
    if extra:
        calc(ws.cell(r,13,f'=SUM(M{r0}:M{r-1})'),"0",True)
        calc(ws.cell(r,14,f'=IF(M{r}=0,0,SUMPRODUCT(M{r0}:M{r-1},N{r0}:N{r-1})/M{r})'),"0.0",True)
        for col in "OP": calc(ws[f"{col}{r}"],MONEY0,True); ws[f"{col}{r}"].value=f"=SUM({col}{r0}:{col}{r-1})"
    ws.cell(r0-1,11,"Nights available").font=font(color="8C959F",size=8); ws.cell(r0-1,12,"Rent").font=font(color="8C959F",size=8)
month_rows(d,10,"$B$4","$D$4")
# per-property table
header(d,26,["Property","Nights booked","Occupancy","ADR","RevPAR","Payout","Cleaning costs","Expenses","Profit"])
d.cell(25,1,"BY PROPERTY").font=font(bold=True,color=ACC)
for i in range(NP):
    r=27+i; d.cell(r,1,f'=IF(Properties!A{P0+i}="","",Properties!A{P0+i})'); calc(d.cell(r,1))
    sel=f'$A{r}'
    M=metrics(YS,YE,sel)
    calc(d.cell(r,2,f'=IF({sel}="","",{M["nights"]})'),"0")
    d.cell(r,11,f'=IF({sel}="","",{M["avail"]})'); d.cell(r,11).font=font(color="8C959F",size=8); d.cell(r,11).number_format="0"
    d.cell(r,12,f'=IF({sel}="","",{M["rent"]})'); d.cell(r,12).font=font(color="8C959F",size=8); d.cell(r,12).number_format=MONEY0
    calc(d.cell(r,3,f'=IF({sel}="","",IF(K{r}=0,0,B{r}/K{r}))'),PCT0)
    calc(d.cell(r,4,f'=IF({sel}="","",IF(B{r}=0,0,L{r}/B{r}))'),MONEY0)
    calc(d.cell(r,5,f'=IF({sel}="","",IF(K{r}=0,0,L{r}/K{r}))'),MONEY0)
    calc(d.cell(r,6,f'=IF({sel}="","",{M["payout"]})'),MONEY0)
    calc(d.cell(r,7,f'=IF({sel}="","",{M["clean"]})'),MONEY0)
    calc(d.cell(r,8,f'=IF({sel}="","",SUMIFS({XR("F")},{XR("A")},">="&{YS},{XR("A")},"<"&{YE},{XR("B")},{sel}))'),MONEY0)
    calc(d.cell(r,9,f'=IF({sel}="","",F{r}-G{r}-H{r})'),MONEY0)
r=27+NP
d.cell(r,1,"Shared costs (no property)").font=font(italic=True)
calc(d.cell(r,8,f'=H22-SUM(H27:H{r-1})'),MONEY0); calc(d.cell(r,9,f'=-H{r}'),MONEY0)
d.cell(r+1,1,"Note: Tax set-aside is a planning estimate, not tax advice. Stays that cross months are split by nights.").font=font(italic=True,size=9,color=GREY)
d.cell(r+2,1,"Year total profit (row 22) = sum of properties minus shared costs.").font=font(italic=True,size=9,color=GREY)
d.conditional_formatting.add("I10:I22",CellIsRule(operator="lessThan",formula=['0'],font=Font(name=F,color="9C1C28",bold=True)))
d.conditional_formatting.add(f"I27:I{r}",CellIsRule(operator="lessThan",formula=['0'],font=Font(name=F,color="9C1C28",bold=True)))
ch=BarChart(); ch.type="col"; ch.title="Payout vs profit by month"; ch.height=8; ch.width=18
ch.add_data(Reference(d,min_col=6,min_row=9,max_row=21),titles_from_data=True)
ch.add_data(Reference(d,min_col=9,min_row=9,max_row=21),titles_from_data=True)
ch.set_categories(Reference(d,min_col=1,min_row=10,max_row=21)); ch.y_axis.numFmt='$#,##0'
d.add_chart(ch,"N4")
d.column_dimensions["K"].width=9; d.column_dimensions["L"].width=9

# ---------------- MONTHLY SUMMARY
ms_=sheet("Monthly Summary")
title(ms_,"Monthly Summary","Pick a property (or All properties) and a year.")
colw(ms_,[16,11,11,11,11,13,13,13,13,4,9,9,10,10,12,12])
ms_["A4"]="Property"; ms_["A4"].font=font(bold=True); inp(ms_["B4"]); ms_["B4"].value="All properties"; ms_.merge_cells("B4:D4")
ms_["F4"]="Year"; ms_["F4"].font=font(bold=True); inp(ms_["G4"],"0"); ms_["G4"].value="=Dashboard!B4"
dv(ms_,"B4","Settings!$L$4:$L$24")
header(ms_,6,["Month","Nights booked","Occupancy","ADR","RevPAR","Payout","Cleaning costs","Expenses","Profit"])
header(ms_,6,["Stays (by check-in)","Avg stay (nights)","Cleaning fees collected","Platform fees"],start=13)
month_rows(ms_,7,"$G$4","$B$4",extra=True)
ms_.cell(21,1,"Expenses: with one property picked, only that property's expenses count; shared costs appear under All properties. Stays and average stay count by check-in month.").font=font(italic=True,size=9,color=GREY)
ms_.conditional_formatting.add("I7:I19",CellIsRule(operator="lessThan",formula=['0'],font=Font(name=F,color="9C1C28",bold=True)))
ch=BarChart(); ch.type="col"; ch.title="Occupancy by month"; ch.height=7; ch.width=18; ch.y_axis.numFmt='0%'
ch.add_data(Reference(ms_,min_col=3,min_row=6,max_row=18),titles_from_data=True)
ch.set_categories(Reference(ms_,min_col=1,min_row=7,max_row=18)); ch.legend=None
ms_.add_chart(ch,"A23")

order=["Start Here","Dashboard","Monthly Summary","Properties","Bookings","Turnovers","Expenses","Supplies","Pricing Calculator","Settings"]
wb._sheets=[wb[n] for n in order]
wb.active=0
wb.save(sys.argv[1] if len(sys.argv)>1 else "Airbnb-Host-Business-Manager.xlsx")
print("saved")
