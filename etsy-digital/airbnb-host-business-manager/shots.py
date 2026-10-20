"""Listing image HTML from the recalculated demo.xlsx (same layout system as _engine/shots.py)."""
import datetime, html
from openpyxl import load_workbook
wb=load_workbook('demo.xlsx',data_only=True)
ACC="#1E6B64"; LIGHT="#E3F1EF"; ACC2="#8CC7BE"; NAME="Airbnb Host Business Manager"
def fmt(v,kind=None):
    if v is None or v=="": return ""
    if isinstance(v,(datetime.datetime,datetime.date)): return v.strftime("%b %-d, %Y") if kind!="d" else v.strftime("%b %-d")
    if kind=="$": return "-" if v==0 else (("−$" if v<0 else "$")+f"{abs(v):,.2f}")
    if kind=="$0": return "-" if v==0 else (("−$" if v<0 else "$")+f"{abs(v):,.0f}")
    if kind=="%": return f"{v*100:.0f}%"
    if isinstance(v,float): return f"{v:g}"
    return html.escape(str(v))
TABS=["Start Here","Dashboard","Monthly Summary","Properties","Bookings","Turnovers","Expenses","Supplies","Pricing Calculator","Settings"]
def win(tab,body,w=1700):
    t="".join(f'<span class="tab {"on" if x==tab else ""}">{x}</span>' for x in TABS)
    return f'<div class="win" style="width:{w}px"><div class="bar"><i></i><i></i><i></i><span>{NAME}</span></div><div class="sheet">{body}</div><div class="tabs">{t}</div></div>'
CSS=f"""*{{box-sizing:border-box}}body{{margin:0;width:2000px;height:1600px;font-family:Inter,Arial,sans-serif;background:linear-gradient(135deg,{LIGHT},#F8F5EE);display:flex;flex-direction:column;align-items:center;justify-content:center;overflow:hidden}}
h1{{font-size:78px;margin:0 0 10px;color:#123c38;letter-spacing:-1px}} h2{{font-size:38px;margin:0 0 40px;color:{ACC};font-weight:600}}
.win{{background:#fff;border-radius:18px;box-shadow:0 30px 80px rgba(20,50,50,.25);overflow:hidden}}
.bar{{background:#f1f3f4;height:48px;display:flex;align-items:center;padding:0 18px;gap:9px;font-size:20px;color:#555}}.bar i{{width:14px;height:14px;border-radius:50%;background:#e0e0e0;display:inline-block}}.bar span{{margin-left:16px}}
.sheet{{padding:26px 30px}} .tabs{{background:#f1f3f4;padding:10px 20px;display:flex;gap:4px}}.tab{{font-size:16px;padding:7px 11px;border-radius:6px;color:#555;white-space:nowrap}}.tab.on{{background:#fff;color:{ACC};font-weight:700;box-shadow:0 1px 3px rgba(0,0,0,.15)}}
table{{border-collapse:collapse;width:100%;font-size:20px}}th{{background:{ACC};color:#fff;padding:11px 10px;font-weight:600;text-align:left}}td{{padding:9px 10px;border-bottom:1px solid #e3e8ea;white-space:nowrap}}td.in{{background:#FFF7D6;color:#1F4EB4}}
td.r,th.r{{text-align:right}} .red{{background:#F8D7DA!important;color:#9C1C28!important;font-weight:700}} .amb{{background:#FFF3CD!important;font-weight:600}} .ok{{color:{ACC};font-weight:700}} tr.tot td{{font-weight:800;border-top:2px solid {ACC}}}
.kpis{{display:flex;gap:18px;margin-bottom:24px}}.kpi{{flex:1;background:{LIGHT};border-radius:14px;padding:16px 20px}}.kpi b{{display:block;font-size:40px;color:{ACC}}}.kpi span{{font-size:19px;color:#57606a}}
.ttl{{font-size:34px;font-weight:800;color:{ACC};margin-bottom:6px}}.sub{{font-size:19px;color:#57606a;margin-bottom:20px;font-style:italic}}
"""
def page(name,inner,head=None,sub=None):
    hh=(f"<h1>{head}</h1>" if head else "")+(f"<h2>{sub}</h2>" if sub else "")
    open(f"img_{name}.html","w").write(f"<html><head><meta charset='utf-8'><style>{CSS}</style></head><body>{hh}{inner}</body></html>")
def tbl(head,rows,right=(),cls=None):
    h="<table><tr>"+"".join(f'<th class="{"r" if i in right else ""}">{x}</th>' for i,x in enumerate(head))+"</tr>"
    for k,row in enumerate(rows):
        trc=' class="tot"' if cls and cls(k,row)=="tot" else ""
        h+=f"<tr{trc}>"
        for i,x in enumerate(row):
            c=[]
            if i in right: c.append("r")
            if isinstance(x,tuple): x,extra=x; c.append(extra)
            h+=f'<td class="{" ".join(c)}">{x}</td>'
        h+="</tr>"
    return h+"</table>"
# ---- Dashboard
d=wb['Dashboard']
kp="".join(f'<div class="kpi"><span>{d[c+"6"].value}</span><b>{fmt(d[c+"7"].value,"%" if c=="E" else "$0")}</b></div>' for c in ("A","C","E","F","H"))
mrows=[r for r in range(10,22) if (d.cell(r,6).value or 0)>0]
mx=max(d.cell(r,6).value for r in mrows)
bars="".join(f'<div style="display:flex;flex-direction:column;align-items:center;gap:6px;flex:1"><div style="display:flex;gap:5px;align-items:flex-end;height:250px"><div style="width:26px;height:{250*d.cell(r,6).value/mx}px;background:{ACC};border-radius:5px 5px 0 0"></div><div style="width:26px;height:{max(0,250*d.cell(r,9).value/mx)}px;background:{ACC2};border-radius:5px 5px 0 0"></div></div><span style="font-size:16px;color:#555">{d.cell(r,1).value[:3]}</span></div>' for r in mrows)
chart=f'<div style="display:flex;gap:6px">{bars}</div><div style="font-size:18px;color:#555;margin-top:8px"><span style="color:{ACC}">■</span> Payout &nbsp; <span style="color:{ACC2}">■</span> Profit</div>'
prow=[r for r in range(27,47) if d.cell(r,1).value]
pt=tbl(["Property","Nights","Occupancy","ADR","RevPAR","Payout","Profit"],
       [[fmt(d.cell(r,1).value),fmt(d.cell(r,2).value),fmt(d.cell(r,3).value,"%"),fmt(d.cell(r,4).value,"$0"),fmt(d.cell(r,5).value,"$0"),fmt(d.cell(r,6).value,"$0"),fmt(d.cell(r,9).value,"$0")] for r in prow],right=range(1,7))
hero=f'<div class="ttl">Dashboard — {d["B4"].value}</div><div class="sub">Every property, every month: updates from your Bookings and Expenses</div><div class="kpis">{kp}</div><div style="display:flex;gap:34px;align-items:flex-start"><div style="flex:1.15">{pt}</div><div style="flex:1">{chart}</div></div>'
page("1_hero",win("Dashboard",hero,1840),NAME,"Bookings · Turnovers · Occupancy · Expenses · Profit — Google Sheets & Excel")
mt=tbl(["Month","Nights","Occupancy","ADR","RevPAR","Payout","Cleaning","Expenses","Profit"],
       [[fmt(d.cell(r,1).value),fmt(d.cell(r,2).value),fmt(d.cell(r,3).value,"%"),fmt(d.cell(r,4).value,"$0"),fmt(d.cell(r,5).value,"$0"),fmt(d.cell(r,6).value,"$0"),fmt(d.cell(r,7).value,"$0"),fmt(d.cell(r,8).value,"$0"),fmt(d.cell(r,9).value,"$0")] for r in list(range(10,22))+[22]],
       right=range(1,9),cls=lambda k,row:"tot" if k==12 else None)
page("2_dashboard",win("Dashboard",f'<div class="ttl">Know your numbers</div><div class="sub">Occupancy, ADR and RevPAR by month. Stays that cross months are split by nights.</div>{mt}',1700),"Occupancy, ADR & RevPAR — done for you",None)
# ---- Bookings
b=wb['Bookings']
last=max(r for r in range(5,1005) if b.cell(r,2).value and b.cell(r,3).value and b.cell(r,3).value.date()<=datetime.date(2026,10,25))
rows=[]
for r in range(last-10,last+1):
    rows.append([fmt(b.cell(r,2).value,"d"),fmt(b.cell(r,4).value),(fmt(b.cell(r,6).value),"in"),fmt(b.cell(r,8).value),fmt(b.cell(r,10).value,"$0"),fmt(b.cell(r,11).value,"$0"),
                 fmt(b.cell(r,14).value,"$0"),fmt(b.cell(r,15).value,"$"),fmt(b.cell(r,17).value,"$"),fmt(b.cell(r,19).value,"$"),(fmt(b.cell(r,21).value),"red" if b.cell(r,21).value=="No" else "")])
bt=tbl(["Check-in","Property","Channel","Nights","Rent","Cleaning fee","Guest total","Platform fee","Payout","Net after cleaning","Paid?"],rows,right=range(3,10))
page("3_bookings",win("Bookings",f'<div class="ttl">Bookings</div><div class="sub">Enter dates, property and channel. Weekend rates, fees, platform cut, payout and net fill in.</div>{bt}',1880),"Every booking, every dollar","Payout and net after cleaning on each stay")
# ---- Turnovers
t=wb['Turnovers']
rows=[]
for r in range(7,19):
    if not t.cell(r,1).value: break
    al=t.cell(r,9).value or ""
    rows.append([fmt(t.cell(r,1).value,"d"),t.cell(r,2).value,fmt(t.cell(r,3).value),fmt(t.cell(r,4).value),fmt(t.cell(r,5).value,"d"),fmt(t.cell(r,6).value),fmt(t.cell(r,7).value),
                 (html.escape(al),"red" if al.startswith("SAME") else "amb" if al.startswith("Next") else ""),(fmt(t.cell(r,10).value),"red" if t.cell(r,10).value=="Unassigned" else ""),(fmt(t.cell(r,11).value),"ok")])
tt=tbl(["Turnover","Day","Property","Checking out","Next check-in","Next guest","Guests","Alert","Cleaner","Done?"],rows,right=(6,))
page("4_turnovers",win("Turnovers",f'<div class="ttl">Turnovers</div><div class="sub">Your cleaning schedule builds itself from Bookings, with same-day turnover alerts</div>{tt}',1880),"Never miss a turnover","Every check-out, the next check-in and who's cleaning")
# ---- Pricing calculator
q=wb['Pricing Calculator']
qi="".join(f'<tr><td>{q.cell(r,1).value}</td><td class="in">{fmt(q.cell(r,2).value)}</td></tr>' for r in range(5,11))
qo="".join(f'<tr style="{"font-weight:800;font-size:26px;color:"+ACC if q.cell(r,1).value=="NET TO YOU" else "font-weight:700" if q.cell(r,1).value.isupper() else ""}"><td>{q.cell(r,1).value.replace("GUEST PAYS","Guest pays").replace("NET TO YOU","Net to you")}</td><td class="r">{fmt(q.cell(r,2).value,"$" if r>13 else None)}</td></tr>' for r in range(12,26))
bi="".join(f'<tr><td>{q.cell(r,4).value}</td><td class="in">{fmt(q.cell(r,5).value,"%" if r==6 else "$0" if r in (5,8) else None)}</td></tr>' for r in range(5,10))
bo="".join(f'<tr style="{"font-weight:800;color:"+ACC if r in (15,16) else ""}"><td>{q.cell(r,4).value.replace("BREAK-EVEN RATE","Break-even rate").replace("RATE FOR TARGET PROFIT","Rate for target profit")}</td><td class="r">{fmt(q.cell(r,5).value,"$" if r in (15,16) else "%" if r==14 else None) if not isinstance(q.cell(r,5).value,float) or r in (14,15,16) else f"{q.cell(r,5).value:.1f}"}</td></tr>' for r in range(12,17))
qb=f'<div style="display:flex;gap:40px"><div style="flex:1"><div class="ttl">Price a stay</div><table><tr><th colspan=2>Stay details</th></tr>{qi}</table><br><table><tr><th colspan=2>What you keep</th></tr>{qo}</table></div><div style="flex:1"><div class="ttl">Break-even rate</div><table><tr><th colspan=2>Your costs</th></tr>{bi}</table><br><table><tr><th colspan=2>Rate you need</th></tr>{bo}</table></div></div>'
page("5_pricing",win("Pricing Calculator",qb.replace("Monthly fixed costs (mortgage, utilities, internet, insurance…)","Monthly fixed costs"),1700),"Price with confidence",None)
# ---- Supplies + Properties
s=wb['Supplies']
rows=[[fmt(s.cell(r,1).value),fmt(s.cell(r,4).value),(fmt(s.cell(r,5).value),"in"),fmt(s.cell(r,6).value),(fmt(s.cell(r,7).value),"red" if s.cell(r,7).value=="Reorder" else "ok"),fmt(s.cell(r,8).value),fmt(s.cell(r,10).value,"$")] for r in range(7,19)]
stt=tbl(["Item","Par","On hand","Reorder at","Status","Order qty","Order cost"],rows,right=(1,2,3,5,6))
p=wb['Properties']
prow=[[fmt(p.cell(r,1).value),fmt(p.cell(r,5).value,"$0"),fmt(p.cell(r,6).value,"$0"),fmt(p.cell(r,7).value,"$0"),(fmt(p.cell(r,13).value),"amb" if p.cell(r,13).value=="Turnover today" else "ok"),fmt(p.cell(r,14).value,"d")] for r in range(5,8)]
ptt=tbl(["Property","Weeknight","Weekend","Cleaning fee","Today","Next check-in"],prow,right=(1,2,3))
sb=f'<div style="display:flex;gap:36px"><div style="flex:1.05"><div class="ttl">Supplies</div><div class="sub">Count after a turnover. Low items turn red. Reorder total: <b>{fmt(s["B4"].value,"$")}</b></div>{stt}</div><div style="flex:1"><div class="ttl">Properties</div><div class="sub">Rates and fees per listing, plus today\'s status</div>{ptt}</div></div>'
page("6_supplies",win("Supplies",sb,1880),"Stocked, priced and ready","Supplies par levels and every property at a glance")
feat=[("📊","Profit dashboard","Payout, profit, occupancy, ADR and RevPAR by month and by property"),("📅","Bookings log","Weekend rates, stay discounts, cleaning & guest fees, platform fee, payout"),
("🧹","Turnover schedule","Next check-outs, same-day alerts, guest counts, cleaner and done status"),("🏠","Multi-property","Up to 20 listings, each with its own rates, fees and cleaner pay"),
("🧮","Pricing calculator","Price any stay and find your break-even nightly rate"),("💸","Expense tracker","Per property or shared, with host categories and receipts"),
("🧻","Supplies tracker","Par levels, reorder alerts, order quantity and cost"),("✅","Google Sheets + Excel","Instant download, works in both. No subscription")]
fb="".join(f'<div style="background:#fff;border-radius:16px;padding:28px;box-shadow:0 10px 30px rgba(20,50,50,.12)"><div style="font-size:46px">{a}</div><div style="font-size:30px;font-weight:800;color:#123c38;margin:8px 0">{b_}</div><div style="font-size:22px;color:#57606a">{c}</div></div>' for a,b_,c in feat)
page("7_included",f'<div style="display:grid;grid-template-columns:repeat(4,400px);gap:26px">{fb}</div><div style="margin-top:34px;font-size:20px;color:#57606a">Not affiliated with Airbnb or Vrbo. Fees and rates are editable examples.</div>',"What's included","10 connected tabs · instant digital download")
print("ok")
